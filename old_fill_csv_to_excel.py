#!/usr/bin/env python3
# -*- coding: utf-8 -*-

"""
現_打鍵前 / 現_打鍵後 の CSV を、テンプレート Excel と同じ形式で展開するツール。

今回の主な仕様:
- 現_打鍵前 / 現_打鍵後 の CSV をファイル名で突合する。
- CSV入力フォルダは複数指定でき、同名CSVは先に指定した有効CSVを優先する。
- 優先フォルダ側のCSVが存在し「無効な識別」を含まない場合は、ヘッダーのみでも採用する。
- 「無効な識別」を含む場合のみ、後続フォルダの同名CSVを採用する。
- 片側にしかない CSV は、存在する側のヘッダーを使い、欠損側は0件扱いにする。
- CSV 内に「無効な識別」がある場合は、そのCSVを0件扱いにする。
- 現_打鍵前件数=現_打鍵後件数=0 の表はテーブルシートへ出力せず、テーブル検証 のみに出力する。
- テーブル検証 には CRUD / 出力番号 / 出力Sheet / テーブル名 / 現・新の件数チェック / 差分チェック / 備考 を出力する。
- テーブル検証 の件数・差異・チェック項目は、固定値ではなく Excel 計算式で出力する。
- CRUD が空の テーブル検証 行は既定で非表示にする。
- 既定で 1 sheet につき 5 テーブルを出力し、超過時は テーブル1, テーブル2... を作成する。
- テーブル見出しは「1.テーブル名」のように番号を付ける。見出しの .csv は削除する。
- CSV 列数が 15 列を超える場合はテンプレート末尾列のスタイルを右方向へ拡張する。
- CSV 列数が 15 列未満の場合は、余った列の内容・スタイル・罫線を消す。
- フォントは全セル Yu Gothic、通常行高 18.75、ヘッダー行高 25.5 に統一し、上下中央にする。
- 比較式・条件付き書式を保持する。
"""

import argparse
import csv
import json
import io
import logging
import os
import posixpath
import re
import shutil
import tempfile
import uuid
import warnings
import zipfile
import xml.etree.ElementTree as ET
from copy import copy
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path, PurePosixPath

from openpyxl import load_workbook
from openpyxl.cell.cell import MergedCell
from openpyxl.formatting.rule import CellIsRule
from openpyxl.formula.translate import Translator
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


TEMPLATE_COLS = 17
DATA_START_COL = 3
FONT_NAME = "Yu Gothic"
BODY_ROW_HEIGHT = 18.75
HEADER_ROW_HEIGHT = 25.5

INVALID_MARKER = "無効な識別"
VALID_DELIMITERS = [",", "\t", ";", "|"]

GRAY_FILL = PatternFill(fill_type="solid", fgColor="D9D9D9")
SUMMARY_HEADER_FILL = PatternFill(fill_type="solid", fgColor="5B9BD5")
SUMMARY_C_FILL = PatternFill(fill_type="solid", fgColor="E2F0D9")
SUMMARY_U_FILL = PatternFill(fill_type="solid", fgColor="DDEBF7")
SUMMARY_D_FILL = PatternFill(fill_type="solid", fgColor="FCE4D6")
SUMMARY_SKIP_FILL = PatternFill(fill_type="solid", fgColor="E7E6E6")

TRUE_FILL = PatternFill(fill_type="solid", fgColor="C6EFCE")
FALSE_FILL = PatternFill(fill_type="solid", fgColor="FFC7CE")
FALSE_FONT = Font(name=FONT_NAME, color="9C0006")
WARNING_FILL = PatternFill(fill_type="solid", fgColor="FFFF00")

THIN_SIDE = Side(style="thin", color="000000")
MEDIUM_SIDE = Side(style="medium", color="000000")
NO_SIDE = Side(style=None)

SUMMARY_SHEET_NAME = "テーブル検証"
COMPARE_EXCLUDE_SHEET_NAME = "現新比較対象外"
TRIANGLE_MARK = "△"
CURRENT_NO_UPDATE_NOTE = "現テストで打鍵前後更新なし"

warnings.filterwarnings(
    "ignore",
    message="Conditional Formatting extension is not supported and will be removed",
)


@dataclass
class CsvInfo:
    path: Path | None
    header: list[str] = field(default_factory=list)
    rows: list[list[str]] = field(default_factory=list)
    exists: bool = False
    invalid: bool = False
    repair_notes: list[str] = field(default_factory=list)

    @property
    def count(self) -> int:
        return len(self.rows)


@dataclass
class TableRecord:
    key: str
    display_name: str
    sql_name: str
    before: CsvInfo
    after: CsvInfo
    header: list[str]
    col_count: int
    crud: str
    count_diff: int
    notes: list[str]
    skip_detail: bool
    current_no_update: bool = False
    output_no: int | None = None
    output_sheet: str = ""
    before_current_extraction_condition: str = ""
    before_new_extraction_condition: str = ""
    after_current_extraction_condition: str = ""
    after_new_extraction_condition: str = ""
    detail_refs: dict = field(default_factory=dict)


# ===== CSV =====
def _normalize_csv_rows(rows):
    return [["" if v is None else str(v) for v in row] for row in rows]


def _is_blank_csv_value(value):
    return str(value or "").strip(" ") == ""


def _is_blank_csv_row(row):
    return all(_is_blank_csv_value(value) for value in row)


def _repair_broken_csv_data_rows(rows, col_count):
    fixed_rows = []
    blank_row_count = 0
    merged_row_count = 0
    for row in rows:
        row = list(row)
        if _is_blank_csv_row(row):
            blank_row_count += 1
            continue
        padded = row + [""] * max(0, col_count - len(row))
        nonblank_indexes = [idx for idx, value in enumerate(padded) if not _is_blank_csv_value(value)]
        if len(nonblank_indexes) == 1 and fixed_rows and col_count > 5:
            idx = nonblank_indexes[0]
            if len(fixed_rows[-1]) <= idx:
                fixed_rows[-1].extend([""] * (idx + 1 - len(fixed_rows[-1])))
            previous = fixed_rows[-1][idx]
            current = padded[idx]
            fixed_rows[-1][idx] = f"{previous}\n{current}" if previous else current
            merged_row_count += 1
            continue
        fixed_rows.append(padded[:col_count] if col_count > 0 else row)
    return fixed_rows, blank_row_count, merged_row_count


def _infer_delimiter(text):
    sample_lines = []
    for line in text.splitlines():
        if line.strip():
            sample_lines.append(line)
        if len(sample_lines) >= 20:
            break
    sample = "\n".join(sample_lines)
    if sample:
        try:
            dialect = csv.Sniffer().sniff(sample, delimiters="".join(VALID_DELIMITERS))
            delimiter = getattr(dialect, "delimiter", None)
            if isinstance(delimiter, str) and len(delimiter) == 1 and delimiter in VALID_DELIMITERS:
                return delimiter
        except (csv.Error, TypeError, ValueError):
            pass
    best_delimiter = ","
    best_count = -1
    for delimiter in VALID_DELIMITERS:
        count = sum(line.count(delimiter) for line in sample_lines)
        if count > best_count:
            best_count = count
            best_delimiter = delimiter
    return best_delimiter


def _read_csv_from_text(text):
    delimiter = _infer_delimiter(text)
    try:
        return list(csv.reader(io.StringIO(text), delimiter=delimiter))
    except (csv.Error, ValueError):
        return list(csv.reader(io.StringIO(text), delimiter=","))


def _decode_csv_text(path: Path):
    raw = path.read_bytes()
    for enc in ("utf-8-sig", "cp932", "shift_jis", "utf-8", "gb18030"):
        try:
            return raw.decode(enc), enc
        except UnicodeDecodeError:
            continue
    return raw.decode("cp932", errors="replace"), "cp932-replace"


def read_csv_info(path: Path | None) -> CsvInfo:
    if path is None:
        return CsvInfo(path=None, exists=False)
    text, _enc = _decode_csv_text(path)
    rows = _normalize_csv_rows(_read_csv_from_text(text))
    first_nonblank_idx = next((idx for idx, row in enumerate(rows) if not _is_blank_csv_row(row)), None)
    blank_before_header_count = len(rows) if first_nonblank_idx is None else sum(1 for row in rows[:first_nonblank_idx] if _is_blank_csv_row(row))
    if first_nonblank_idx is None:
        rows = []
        header = []
        data_source_rows = []
    else:
        header = rows[first_nonblank_idx]
        data_source_rows = rows[first_nonblank_idx + 1:]
    invalid = INVALID_MARKER in text
    data_col_count = max(len(header), *(len(r) for r in data_source_rows), 0)
    repair_notes = []
    if invalid:
        data = []
    else:
        data, blank_data_row_count, merged_row_count = _repair_broken_csv_data_rows(data_source_rows, data_col_count)
        blank_row_count = blank_before_header_count + blank_data_row_count
        if blank_row_count:
            repair_notes.append(f"空白行を{blank_row_count}行無視")
        if merged_row_count:
            repair_notes.append(f"改行により分割された行を{merged_row_count}行結合")
    return CsvInfo(path=path, header=header, rows=data, exists=True, invalid=invalid, repair_notes=repair_notes)


def is_effective_csv_info(info: CsvInfo) -> bool:
    """優先採用できるCSVか判定する。

    条件:
    - CSVが存在する
    - 「無効な識別」を含まない

    ヘッダーのみでも採用し、無効な識別を含む場合だけ後続候補へフォールバックする。
    """
    return bool(info.exists and not info.invalid)


def csv_files(folder):
    folders = folder if isinstance(folder, (list, tuple)) else [folder]
    result = {}
    first_bad = {}
    for f in folders:
        f = Path(f)
        if not f.exists():
            raise FileNotFoundError(f"フォルダが存在しません: {f}")
        for p in f.iterdir():
            if not (p.is_file() and p.suffix.lower() == ".csv"):
                continue
            key = p.name.lower()
            # 有効CSVが既に採用済みなら、低優先フォルダは見ない。
            if key in result:
                continue

            info = read_csv_info(p)
            if is_effective_csv_info(info):
                result[key] = p
            else:
                # 無効な識別などの場合は、後続フォルダの同名CSVを探す。
                # ただし全候補が無効だった場合に備え、最初の無効候補は保持する。
                first_bad.setdefault(key, p)

    for key, p in first_bad.items():
        result.setdefault(key, p)
    return result


# ===== template / style =====
def table_name_from_file(path: Path, mode: str):
    return path.name if mode == "filename" else path.stem


def row_text(ws, row):
    values = []
    max_c = min(max(ws.max_column, TEMPLATE_COLS), TEMPLATE_COLS)
    for c in range(1, max_c + 1):
        v = ws.cell(row, c).value
        if v is not None:
            values.append(str(v))
    return " ".join(values)


def find_after_sql_template_row(tws):
    max_r = min(tws.max_row, 25)
    for r in range(14, max_r + 1):
        text = row_text(tws, r)
        upper = text.upper()
        if ("取得SQL" in text or "SELECT" in upper) and ("新" not in text):
            return r
        if "SELECT" in upper and ("WHERE" in upper or "抽出条件" in text or "テーブル" in text):
            return r
    text16 = row_text(tws, 16)
    if "更新" in text16 and "件数" in text16:
        return 17
    return 16


def template_layout(tws):
    after_sql = find_after_sql_template_row(tws)
    def has_new_sql_row(start, end):
        for r in range(max(1, start), min(tws.max_row, end) + 1):
            text = row_text(tws, r)
            if "取得SQL" in text and "新" in text:
                return True
        return False

    before_has_new_sql = has_new_sql_row(3, 6)
    after_has_new_sql = has_new_sql_row(after_sql, after_sql + 3)
    before_offset = 1 if before_has_new_sql else 0
    after_offset = 1 if after_has_new_sql else 0
    return {
        "after_sql": after_sql,
        "before_sql_new": 4 if before_has_new_sql else None,
        "before_header": 4 + before_offset,
        "before_data": 5 + before_offset,
        "before_new": 8 + before_offset,
        "before_compare": 11 + before_offset,
        "after_sql_new": after_sql + 1 if after_has_new_sql else None,
        "after_header": after_sql + 1 + after_offset,
        "after_data": after_sql + 2 + after_offset,
        "after_new": after_sql + 5 + after_offset,
        "after_compare": after_sql + 8 + after_offset,
        "count_title": after_sql + 11 + after_offset,
        "before_count": after_sql + 12 + after_offset,
        "after_count": after_sql + 13 + after_offset,
        "count_diff": after_sql + 14 + after_offset,
        "pre_cross_title": after_sql + 15 + after_offset,
        "cross_title": after_sql + 16 + after_offset,
        "cross_current": after_sql + 17 + after_offset,
        "cross_new": after_sql + 20 + after_offset,
    }


def is_update_count_template_row(tws, row):
    text = row_text(tws, row)
    return "更新" in text and "件数" in text


def is_formula(value):
    return isinstance(value, str) and value.startswith("=")


def _copy_style_component(cell, attr_name, default_value):
    try:
        value = getattr(cell, attr_name)
        return copy(value) if value is not None else copy(default_value)
    except (IndexError, KeyError, TypeError, ValueError, AttributeError):
        return copy(default_value)


def force_font_name(cell, font_name=FONT_NAME):
    if isinstance(cell, MergedCell):
        return
    f = _copy_style_component(cell, "font", Font())
    f.name = font_name
    cell.font = f
    a = _copy_style_component(cell, "alignment", Alignment())
    a.vertical = "center"
    cell.alignment = a


def copy_cell_format(src, dst):
    if isinstance(dst, MergedCell):
        return
    # workbookを跨ぐ場合に _style を直接コピーすると、openpyxl の style index が壊れて
    # cell.border 参照時に IndexError になることがあるため、各style要素を値としてコピーする。
    dst.font = _copy_style_component(src, "font", Font(name=FONT_NAME))
    dst.fill = _copy_style_component(src, "fill", PatternFill(fill_type=None))
    dst.border = _copy_style_component(src, "border", Border())
    dst.alignment = _copy_style_component(src, "alignment", Alignment(vertical="center"))
    try:
        dst.number_format = src.number_format or "General"
    except (IndexError, KeyError, TypeError, ValueError, AttributeError):
        dst.number_format = "General"
    try:
        dst.protection = copy(src.protection)
    except (IndexError, KeyError, TypeError, ValueError, AttributeError):
        pass
    try:
        if src.comment:
            dst.comment = copy(src.comment)
    except (IndexError, KeyError, TypeError, ValueError, AttributeError):
        pass
    force_font_name(dst)


def clear_cell_format(cell):
    if isinstance(cell, MergedCell):
        return
    cell.value = None
    cell.border = Border()
    cell.fill = PatternFill(fill_type=None)
    cell.font = Font(name=FONT_NAME)
    cell.alignment = Alignment(vertical="center")
    cell.number_format = "General"


def clear_extra_columns(ws, row, first_extra_col, last_col):
    if first_extra_col > last_col:
        return
    for c in range(first_extra_col, last_col + 1):
        clear_cell_format(ws.cell(row, c))


def template_cell_for_col(tws, template_row, dest_col):
    return tws.cell(template_row, dest_col if dest_col <= TEMPLATE_COLS else TEMPLATE_COLS)


def translate_formula(formula, origin, target):
    try:
        return Translator(formula, origin=origin).translate_formula(target)
    except Exception:
        return formula


def translate_cf_formula(formula, origin, target):
    if not isinstance(formula, str):
        return formula
    if not re.search(r"\$?[A-Za-z]{1,3}\$?\d+", formula):
        return formula
    had_equal = formula.startswith("=")
    expr = formula if had_equal else "=" + formula
    try:
        translated = Translator(expr, origin=origin).translate_formula(target)
        return translated if had_equal else translated[1:]
    except Exception:
        return formula


def copy_conditional_formatting_for_row(tws, ows, template_row, out_row, block_last_col):
    cf_rules = getattr(tws.conditional_formatting, "_cf_rules", None)
    if not cf_rules:
        return
    for cf_obj, rules in cf_rules.items():
        for src_range in cf_obj.sqref.ranges:
            if not (src_range.min_row <= template_row <= src_range.max_row):
                continue
            dest_min_col = src_range.min_col
            dest_max_col = block_last_col if src_range.max_col >= TEMPLATE_COLS else min(src_range.max_col, block_last_col)
            if dest_min_col > block_last_col or dest_min_col > dest_max_col:
                continue
            dest_ref = f"{get_column_letter(dest_min_col)}{out_row}:{get_column_letter(dest_max_col)}{out_row}"
            if dest_min_col == dest_max_col:
                dest_ref = f"{get_column_letter(dest_min_col)}{out_row}"
            src_anchor = f"{get_column_letter(src_range.min_col)}{template_row}"
            dst_anchor = f"{get_column_letter(dest_min_col)}{out_row}"
            for rule in rules:
                new_rule = copy(rule)
                if getattr(rule, "formula", None):
                    new_rule.formula = [translate_cf_formula(f, src_anchor, dst_anchor) for f in rule.formula]
                ows.conditional_formatting.add(dest_ref, new_rule)


def copy_template_row_style(tws, ows, template_row, out_row, block_last_col):
    for c in range(1, block_last_col + 1):
        copy_cell_format(template_cell_for_col(tws, template_row, c), ows.cell(out_row, c))
    copy_conditional_formatting_for_row(tws, ows, template_row, out_row, block_last_col)


def set_cell_border(cell, left=None, right=None, top=None, bottom=None):
    if isinstance(cell, MergedCell):
        return
    try:
        b = copy(cell.border)
    except (IndexError, KeyError, TypeError, ValueError, AttributeError):
        b = Border()
    cell.border = Border(
        left=left if left is not None else b.left,
        right=right if right is not None else b.right,
        top=top if top is not None else b.top,
        bottom=bottom if bottom is not None else b.bottom,
        diagonal=b.diagonal,
        diagonal_direction=b.diagonal_direction,
        diagonalUp=b.diagonalUp,
        diagonalDown=b.diagonalDown,
        outline=b.outline,
        vertical=b.vertical,
        horizontal=b.horizontal,
    )


def apply_standard_grid(ws, row_start, row_count, col_start, col_end):
    if row_count <= 0 or col_end < col_start:
        return
    for r in range(row_start, row_start + row_count):
        for c in range(col_start, col_end + 1):
            set_cell_border(ws.cell(r, c), left=THIN_SIDE, right=THIN_SIDE, top=THIN_SIDE, bottom=THIN_SIDE)


def apply_horizontal_border(ws, row, col_start, col_end, side, where="top"):
    if col_end < col_start:
        return
    for c in range(col_start, col_end + 1):
        if where == "bottom":
            set_cell_border(ws.cell(row, c), bottom=side)
        else:
            set_cell_border(ws.cell(row, c), top=side)


def clear_horizontal_border(ws, row, col_start, col_end, where="top"):
    if col_end < col_start:
        return
    for c in range(col_start, col_end + 1):
        if where == "bottom":
            set_cell_border(ws.cell(row, c), bottom=NO_SIDE)
        else:
            set_cell_border(ws.cell(row, c), top=NO_SIDE)


def apply_title_row_fill(tws, ows, out_row, block_last_col):
    src_fill = copy(tws.cell(1, 1).fill)
    for c in range(1, block_last_col + 1):
        ows.cell(out_row, c).fill = copy(src_fill)


def write_static_from_template(tws, ows, template_row, out_row, block_last_col):
    copy_template_row_style(tws, ows, template_row, out_row, block_last_col)
    for c in range(1, min(TEMPLATE_COLS, block_last_col) + 1):
        src = tws.cell(template_row, c)
        if src.value is not None:
            dst = ows.cell(out_row, c)
            dst.value = translate_formula(src.value, src.coordinate, dst.coordinate) if is_formula(src.value) else src.value
            force_font_name(dst)


def write_static_data_limited_from_template(tws, ows, template_row, out_row, data_last_col, block_last_col):
    write_static_from_template(tws, ows, template_row, out_row, block_last_col)
    clear_extra_columns(ows, out_row, data_last_col + 1, block_last_col)


def merge_b_label(ows, row_start, row_count, label):
    if row_count <= 0:
        return
    ows.cell(row_start, 2).value = label
    force_font_name(ows.cell(row_start, 2))
    if row_count > 1:
        ows.merge_cells(start_row=row_start, start_column=2, end_row=row_start + row_count - 1, end_column=2)


def fill_gray_row(ows, row, col_start, col_end):
    for c in range(col_start, col_end + 1):
        ows.cell(row, c).fill = copy(GRAY_FILL)


def build_compare_formula(row_a, row_b, col, table_aliases=None, item_name=None, compare_exclusions=None, blank_skip=True):
    letter = get_column_letter(col)
    exact_expr = f"EXACT({letter}{row_a},{letter}{row_b})"
    if is_compare_excluded(table_aliases or set(), item_name, compare_exclusions or set()):
        body = f'IF({exact_expr},TRUE,"{TRIANGLE_MARK}")'
    else:
        body = exact_expr
    if blank_skip:
        return f'=IF(AND({letter}{row_a}="",{letter}{row_b}=""),"",{body})'
    return f"={body}"


def build_result_formula(row, first_col, last_col):
    return f'=IF(COUNTIF({get_column_letter(first_col)}{row}:{get_column_letter(last_col)}{row},FALSE)>0,"差分あり","差分なし")'


def prepare_columns(tws, ows, block_last_col):
    for c in range(1, block_last_col + 1):
        letter = get_column_letter(c)
        src_letter = letter if c <= TEMPLATE_COLS else get_column_letter(TEMPLATE_COLS)
        ows.column_dimensions[letter].width = tws.column_dimensions[src_letter].width or 12


def normalize_matrix(rows, row_count, col_count):
    out = []
    for i in range(row_count):
        src = rows[i] if i < len(rows) else []
        out.append(list(src[:col_count]) + [""] * max(0, col_count - len(src)))
    return out


def add_compare_conditional_formatting(ws, row_start, row_count, data_last_col):
    if row_count <= 0 or data_last_col < DATA_START_COL:
        return
    ref = f"{get_column_letter(DATA_START_COL)}{row_start}:{get_column_letter(data_last_col)}{row_start + row_count - 1}"
    ws.conditional_formatting.add(ref, CellIsRule(operator="equal", formula=["TRUE"], fill=TRUE_FILL))
    ws.conditional_formatting.add(ref, CellIsRule(operator="equal", formula=["FALSE"], fill=FALSE_FILL, font=FALSE_FONT))
    ws.conditional_formatting.add(ref, CellIsRule(operator="equal", formula=[f'"{TRIANGLE_MARK}"'], fill=WARNING_FILL))


def write_header_row(tws, ows, template_row, out_row, header, col_count, block_last_col, header_rows):
    data_last_col = DATA_START_COL + col_count - 1
    copy_template_row_style(tws, ows, template_row, out_row, block_last_col)
    for idx in range(col_count):
        cell = ows.cell(out_row, DATA_START_COL + idx)
        cell.value = header[idx] if idx < len(header) else ""
        force_font_name(cell)
    clear_extra_columns(ows, out_row, data_last_col + 1, block_last_col)
    apply_standard_grid(ows, out_row, 1, DATA_START_COL, data_last_col)
    header_rows.add(out_row)


def write_data_rows(tws, ows, template_row, start_row, rows, display_count, col_count, block_last_col, gray_if_empty=False):
    data_last_col = DATA_START_COL + col_count - 1
    matrix = normalize_matrix(rows, display_count, col_count)
    for i in range(display_count):
        r = start_row + i
        copy_template_row_style(tws, ows, template_row, r, block_last_col)
        for idx in range(col_count):
            cell = ows.cell(r, DATA_START_COL + idx)
            cell.value = matrix[i][idx]
            force_font_name(cell)
        clear_extra_columns(ows, r, data_last_col + 1, block_last_col)
        if gray_if_empty:
            fill_gray_row(ows, r, DATA_START_COL, data_last_col)
    apply_standard_grid(ows, start_row, display_count, DATA_START_COL, data_last_col)


def write_blank_rows(tws, ows, template_row, start_row, display_count, col_count, block_last_col):
    data_last_col = DATA_START_COL + col_count - 1
    for i in range(display_count):
        r = start_row + i
        copy_template_row_style(tws, ows, template_row, r, block_last_col)
        for idx in range(col_count):
            cell = ows.cell(r, DATA_START_COL + idx)
            cell.value = None
            force_font_name(cell)
        clear_extra_columns(ows, r, data_last_col + 1, block_last_col)
    apply_standard_grid(ows, start_row, display_count, DATA_START_COL, data_last_col)


def write_vertical_compare_rows(tws, ows, template_row, start_row, count, row_a_start, row_b_start, col_count, block_last_col, table_aliases=None, header=None, compare_exclusions=None):
    data_last_col = DATA_START_COL + col_count - 1
    table_aliases = table_aliases or set()
    header = header or []
    compare_exclusions = compare_exclusions or set()
    for i in range(count):
        r = start_row + i
        copy_template_row_style(tws, ows, template_row, r, block_last_col)
        ows.cell(r, 2).value = build_result_formula(r, DATA_START_COL, data_last_col)
        force_font_name(ows.cell(r, 2))
        for idx in range(col_count):
            c = DATA_START_COL + idx
            item_name = header[idx] if idx < len(header) else ""
            cell = ows.cell(r, c)
            cell.value = build_compare_formula(row_a_start + i, row_b_start + i, c, table_aliases, item_name, compare_exclusions, blank_skip=True)
            force_font_name(cell)
        clear_extra_columns(ows, r, data_last_col + 1, block_last_col)
    apply_standard_grid(ows, start_row, count, DATA_START_COL, data_last_col)
    add_compare_conditional_formatting(ows, start_row, count, data_last_col)


def write_cross_compare_rows(tws, ows, template_row, start_row, count, row_a_start, row_b_start, label, col_count, block_last_col, table_aliases=None, header=None, compare_exclusions=None):
    data_last_col = DATA_START_COL + col_count - 1
    table_aliases = table_aliases or set()
    header = header or []
    compare_exclusions = compare_exclusions or set()
    for i in range(count):
        r = start_row + i
        copy_template_row_style(tws, ows, template_row, r, block_last_col)
        for idx in range(col_count):
            c = DATA_START_COL + idx
            item_name = header[idx] if idx < len(header) else ""
            cell = ows.cell(r, c)
            cell.value = build_compare_formula(row_a_start + i, row_b_start + i, c, table_aliases, item_name, compare_exclusions, blank_skip=False)
            force_font_name(cell)
        clear_extra_columns(ows, r, data_last_col + 1, block_last_col)
    apply_standard_grid(ows, start_row, count, DATA_START_COL, data_last_col)
    add_compare_conditional_formatting(ows, start_row, count, data_last_col)
    merge_b_label(ows, start_row, count, label)


# ===== records / block =====
def build_table_records(before_map, after_map, sql_name_mode):
    keys = sorted(set(before_map.keys()) | set(after_map.keys()))
    records = []
    for key in keys:
        before = read_csv_info(before_map.get(key))
        after = read_csv_info(after_map.get(key))
        display_path = before.path or after.path
        display_name = display_path.stem if display_path else key[:-4] if key.lower().endswith(".csv") else key
        sql_name = table_name_from_file(display_path, sql_name_mode) if display_path else display_name

        notes = []
        if not before.exists:
            notes.append("打鍵前CSVなし。0件扱い")
        if not after.exists:
            notes.append("打鍵後CSVなし。0件扱い")
        if before.invalid:
            notes.append("打鍵前CSVに無効な識別あり。0件扱い")
        if after.invalid:
            notes.append("打鍵後CSVに無効な識別あり。0件扱い")
        if before.exists and not before.invalid and before.count == 0:
            notes.append("打鍵前CSVがヘッダーのみ。0件扱い")
        if after.exists and not after.invalid and after.count == 0:
            notes.append("打鍵後CSVがヘッダーのみ。0件扱い")
        notes.extend(f"打鍵前CSVデータ補正: {note}" for note in before.repair_notes)
        notes.extend(f"打鍵後CSVデータ補正: {note}" for note in after.repair_notes)

        before_header = before.header[:]
        after_header = after.header[:]
        if not before_header:
            before_header = after_header[:]
        if not after_header:
            after_header = before_header[:]
        if not before.header:
            before.header = before_header
        if not after.header:
            after.header = after_header

        max_cols = max(len(before_header), len(after_header), *(len(r) for r in before.rows), *(len(r) for r in after.rows), 1)
        header = before_header if before_header else after_header
        header = list(header[:max_cols]) + [""] * max(0, max_cols - len(header))

        current_no_update = (
            before.exists
            and after.exists
            and not before.invalid
            and not after.invalid
            and before.rows == after.rows
        )

        diff = after.count - before.count
        if current_no_update:
            crud = ""
        elif after.count > before.count:
            crud = "C"
        elif after.count < before.count:
            crud = "D"
        elif before.count != 0:
            crud = "U"
        else:
            crud = ""
        skip_detail = (before.count == 0 and after.count == 0) or current_no_update
        if current_no_update:
            notes.append(CURRENT_NO_UPDATE_NOTE)
        if skip_detail:
            if before.count == 0 and after.count == 0:
                notes.append("打鍵前後とも0件のため明細出力なし")

        records.append(TableRecord(key, display_name, sql_name, before, after, header, max_cols, crud, diff, notes, skip_detail, current_no_update))
    return records


def write_block(tws, ows, start_row, record, layout, header_rows, cf_targets, compare_exclusions=None):
    before_header = record.before.header[:] if record.before.header else record.header[:]
    after_header = record.after.header[:] if record.after.header else record.header[:]
    before_rows = record.before.rows
    after_rows = record.after.rows
    col_count = record.col_count
    data_last_col = DATA_START_COL + col_count - 1
    block_last_col = max(TEMPLATE_COLS, data_last_col)
    table_aliases = make_table_aliases(record.display_name, record.sql_name, record.key)
    prepare_columns(tws, ows, block_last_col)

    before_count = len(before_rows)
    after_count = len(after_rows)
    before_disp = max(1, before_count)
    after_disp = max(1, after_count)
    compare_b_ranges = []
    r = start_row

    write_static_from_template(tws, ows, 1, r, block_last_col)
    ows.cell(r, 1).value = f"{record.output_no}.{record.display_name}" if record.output_no is not None else record.display_name
    force_font_name(ows.cell(r, 1))
    apply_title_row_fill(tws, ows, r, block_last_col)
    header_rows.add(r)
    r += 1

    write_static_from_template(tws, ows, 2, r, block_last_col)
    r += 1
    write_static_from_template(tws, ows, 3, r, block_last_col)
    before_current_condition = record.before_current_extraction_condition
    ows.cell(r, 3).value = f"SELECT * FROM {record.sql_name} WHERE {before_current_condition}" if before_current_condition else f"SELECT * FROM {record.sql_name}"
    force_font_name(ows.cell(r, 3))
    r += 1
    if layout.get("before_sql_new"):
        write_static_from_template(tws, ows, layout["before_sql_new"], r, block_last_col)
        r += 1
    write_header_row(tws, ows, layout["before_header"], r, before_header, col_count, block_last_col, header_rows)
    r += 1

    before_curr_start = r
    write_data_rows(tws, ows, layout["before_data"], before_curr_start, before_rows, before_disp, col_count, block_last_col, gray_if_empty=(before_count == 0))
    merge_b_label(ows, before_curr_start, before_disp, "現")
    r += before_disp

    before_new_start = r
    write_blank_rows(tws, ows, layout["before_new"], before_new_start, before_disp, col_count, block_last_col)
    merge_b_label(ows, before_new_start, before_disp, "新")
    apply_horizontal_border(ows, before_new_start, 2, data_last_col, MEDIUM_SIDE, where="top")
    clear_horizontal_border(ows, before_new_start, data_last_col + 1, block_last_col, where="top")
    r += before_disp

    before_cmp_start = r
    write_vertical_compare_rows(tws, ows, layout["before_compare"], before_cmp_start, before_disp, before_curr_start, before_new_start, col_count, block_last_col, table_aliases, before_header, compare_exclusions)
    cf_targets.append({"kind": "before_compare", "start": before_cmp_start, "end": before_cmp_start + before_disp - 1, "data_last_col": data_last_col})
    compare_b_ranges.append((before_cmp_start, before_cmp_start + before_disp - 1))
    r += before_disp

    before_compare_template_end = layout["before_compare"] + 2
    for tmpl_row in range(before_compare_template_end + 1, layout["after_sql"]):
        if is_update_count_template_row(tws, tmpl_row):
            continue
        write_static_from_template(tws, ows, tmpl_row, r, block_last_col)
        r += 1

    write_static_from_template(tws, ows, layout["after_sql"], r, block_last_col)
    after_current_condition = record.after_current_extraction_condition
    ows.cell(r, 3).value = f"SELECT * FROM {record.sql_name} WHERE {after_current_condition}" if after_current_condition else f"SELECT * FROM {record.sql_name}"
    force_font_name(ows.cell(r, 3))
    r += 1
    if layout.get("after_sql_new"):
        write_static_from_template(tws, ows, layout["after_sql_new"], r, block_last_col)
        r += 1
    write_header_row(tws, ows, layout["after_header"], r, after_header, col_count, block_last_col, header_rows)
    r += 1

    after_curr_start = r
    write_data_rows(tws, ows, layout["after_data"], after_curr_start, after_rows, after_disp, col_count, block_last_col, gray_if_empty=(after_count == 0))
    merge_b_label(ows, after_curr_start, after_disp, "現")
    r += after_disp

    after_new_start = r
    write_blank_rows(tws, ows, layout["after_new"], after_new_start, after_disp, col_count, block_last_col)
    merge_b_label(ows, after_new_start, after_disp, "新")
    apply_horizontal_border(ows, after_new_start, 2, data_last_col, MEDIUM_SIDE, where="top")
    clear_horizontal_border(ows, after_new_start, data_last_col + 1, block_last_col, where="top")
    r += after_disp

    after_cmp_start = r
    write_vertical_compare_rows(tws, ows, layout["after_compare"], after_cmp_start, after_disp, after_curr_start, after_new_start, col_count, block_last_col, table_aliases, after_header, compare_exclusions)
    cf_targets.append({"kind": "after_compare", "start": after_cmp_start, "end": after_cmp_start + after_disp - 1, "data_last_col": data_last_col})
    compare_b_ranges.append((after_cmp_start, after_cmp_start + after_disp - 1))
    r += after_disp

    write_static_from_template(tws, ows, layout["count_title"], r, block_last_col)
    r += 1
    write_static_from_template(tws, ows, layout["before_count"], r, block_last_col)
    ows.cell(r, 3).value = before_count
    force_font_name(ows.cell(r, 3))
    before_count_row = r
    r += 1
    write_static_from_template(tws, ows, layout["after_count"], r, block_last_col)
    ows.cell(r, 3).value = after_count
    force_font_name(ows.cell(r, 3))
    after_count_row = r
    r += 1
    write_static_from_template(tws, ows, layout["count_diff"], r, block_last_col)
    ows.cell(r, 3).value = f"=C{after_count_row}-C{before_count_row}"
    force_font_name(ows.cell(r, 3))
    count_diff = after_count - before_count
    r += 1

    write_static_from_template(tws, ows, layout["pre_cross_title"], r, block_last_col)
    r += 1

    if count_diff == 0:
        write_static_data_limited_from_template(tws, ows, layout["cross_title"], r, data_last_col, block_last_col)
        r += 1
        cross_count = max(1, before_count)
        cross_current_start = r
        write_cross_compare_rows(tws, ows, layout["cross_current"], cross_current_start, cross_count, before_curr_start, after_curr_start, "現", col_count, block_last_col, table_aliases, before_header, compare_exclusions)
        r += cross_count
        cross_new_start = r
        write_cross_compare_rows(tws, ows, layout["cross_new"], cross_new_start, cross_count, before_new_start, after_new_start, "新", col_count, block_last_col, table_aliases, before_header, compare_exclusions)
        apply_horizontal_border(ows, cross_new_start, 2, data_last_col, MEDIUM_SIDE, where="top")
        clear_horizontal_border(ows, cross_new_start, data_last_col + 1, block_last_col, where="top")
        r += cross_count
        cf_targets.append({"kind": "cross_compare", "start": cross_current_start, "end": cross_new_start + cross_count - 1, "data_last_col": data_last_col})
        compare_b_ranges.append((cross_current_start, cross_new_start + cross_count - 1))

    record.detail_refs = {
        "before_count_row": before_count_row,
        "after_count_row": after_count_row,
        "before_new_start": before_new_start,
        "before_new_end": before_new_start + before_disp - 1,
        "after_new_start": after_new_start,
        "after_new_end": after_new_start + after_disp - 1,
        "data_last_col": data_last_col,
        "compare_b_ranges": compare_b_ranges,
    }
    return r, block_last_col


# ===== x14 条件付き書式 =====
def sheet_xml_paths(xlsx_path: Path) -> dict[str, str]:
    with zipfile.ZipFile(xlsx_path, "r") as z:
        wb_root = ET.fromstring(z.read("xl/workbook.xml"))
        ns_main = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
        ns_rel = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"
        sheets_el = wb_root.find(f"{ns_main}sheets")
        sheets = list(sheets_el) if sheets_el is not None else []
        rels_root = ET.fromstring(z.read("xl/_rels/workbook.xml.rels"))
        rel_ns = "{http://schemas.openxmlformats.org/package/2006/relationships}"
        rel_targets = {
            rel.attrib.get("Id"): rel.attrib.get("Target", "")
            for rel in rels_root.findall(f"{rel_ns}Relationship")
        }
        result = {}
        for sheet in sheets:
            name = sheet.attrib.get("name")
            rid = sheet.attrib.get(f"{ns_rel}id")
            target = rel_targets.get(rid, "")
            if not name or not target:
                continue
            if target.startswith("/"):
                result[name] = target.lstrip("/")
            else:
                result[name] = posixpath.normpath(str(PurePosixPath("xl") / target))
        return result


def resolve_sheet_xml_path(xlsx_path: Path, sheet_name=None):
    with zipfile.ZipFile(xlsx_path, "r") as z:
        wb_root = ET.fromstring(z.read("xl/workbook.xml"))
        ns_main = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
        ns_rel = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"
        sheets_el = wb_root.find(f"{ns_main}sheets")
        sheets = list(sheets_el) if sheets_el is not None else []
        target_sheet = None
        for sheet in sheets:
            if sheet_name is None or sheet.attrib.get("name") == sheet_name:
                target_sheet = sheet
                break
        if target_sheet is None:
            return None
        rid = target_sheet.attrib.get(f"{ns_rel}id")
        rels_root = ET.fromstring(z.read("xl/_rels/workbook.xml.rels"))
        rel_ns = "{http://schemas.openxmlformats.org/package/2006/relationships}"
        for rel in rels_root.findall(f"{rel_ns}Relationship"):
            if rel.attrib.get("Id") == rid:
                target = rel.attrib.get("Target", "")
                if target.startswith("/"):
                    return target.lstrip("/")
                return str(PurePosixPath("xl") / target)
    raise RuntimeError("対象 worksheet XML を特定できません。")


def parse_sqref_rows(sqref):
    first = (sqref or "").split()[0]
    rows = [int(x) for x in re.findall(r"\$?[A-Za-z]{1,3}\$?(\d+)", first)]
    if not rows:
        return None, None
    return min(rows), max(rows)


def extract_x14_compare_cf_templates(template_path: Path, layout, sheet_name=None):
    result = {}
    try:
        sheet_xml_path = resolve_sheet_xml_path(template_path, sheet_name)
        if not sheet_xml_path:
            return result
        with zipfile.ZipFile(template_path, "r") as z:
            xml = z.read(sheet_xml_path).decode("utf-8")
    except Exception:
        return result
    blocks = re.findall(r"<x14:conditionalFormatting\b[\s\S]*?</x14:conditionalFormatting>", xml)
    for block in blocks:
        m = re.search(r"<xm:sqref>(.*?)</xm:sqref>", block)
        if not m:
            continue
        min_row, max_row = parse_sqref_rows(m.group(1))
        if min_row is None:
            continue
        item = {"block": block, "src_min_row": min_row, "src_max_row": max_row}
        if min_row <= layout["before_compare"] <= max_row:
            result.setdefault("before_compare", item)
        if min_row <= layout["after_compare"] <= max_row:
            result.setdefault("after_compare", item)
        if min_row <= layout["cross_current"] <= max_row or min_row <= layout["cross_new"] <= max_row:
            result.setdefault("cross_compare", item)
    return result


def render_x14_cf_block(template_item, dest_start, dest_end, priority_start):
    block = template_item["block"]
    src_anchor = f"A{template_item['src_min_row']}"
    dst_anchor = f"A{dest_start}"
    priority = priority_start
    block = block.replace("<x14:conditionalFormattings>", "").replace("</x14:conditionalFormattings>", "")

    def repl_formula(m):
        return f"<xm:f>{translate_cf_formula(m.group(1), src_anchor, dst_anchor)}</xm:f>"

    def repl_rule(m):
        nonlocal priority
        rule = m.group(0)
        rule = re.sub(r'priority="\d+"', f'priority="{priority}"', rule, count=1)
        priority += 1
        rule = re.sub(r'id="\{[^\"]+\}"', f'id="{{{str(uuid.uuid4()).upper()}}}"', rule, count=1)
        rule = re.sub(r"<xm:f>(.*?)</xm:f>", repl_formula, rule)
        return rule

    block = re.sub(r"<x14:cfRule\b[\s\S]*?</x14:cfRule>", repl_rule, block)
    block = re.sub(r"<xm:sqref>.*?</xm:sqref>", f"<xm:sqref>A{dest_start}:XFD{dest_end}</xm:sqref>", block, count=1)
    return block, priority


def inject_x14_conditional_formatting_for_sheet(xlsx_path: Path, template_path: Path, sheet_name, cf_targets, layout, template_sheet_name=None):
    if not cf_targets:
        return False
    templates = extract_x14_compare_cf_templates(template_path, layout, template_sheet_name)
    if not templates:
        return False
    blocks = []
    priority = 3
    for target in cf_targets:
        tmpl = templates.get(target["kind"])
        if not tmpl:
            continue
        block, priority = render_x14_cf_block(tmpl, target["start"], target["end"], priority)
        blocks.append(block)
    if not blocks:
        return False
    ext_xml = ('<ext uri="{78C0D931-6437-407d-A8EE-F0AAD7539E65}" '
               'xmlns:x14="http://schemas.microsoft.com/office/spreadsheetml/2009/9/main">'
               '<x14:conditionalFormattings>' + "".join(blocks) + '</x14:conditionalFormattings></ext>')
    sheet_xml_path = resolve_sheet_xml_path(xlsx_path, sheet_name)
    if not sheet_xml_path:
        return False
    tmp_dir = xlsx_path.parent if xlsx_path.parent.exists() else Path(".")
    tmp = tempfile.NamedTemporaryFile(prefix=f".{xlsx_path.stem}_", suffix=".tmp.xlsx", dir=tmp_dir, delete=False)
    tmp_name = tmp.name
    tmp.close()
    try:
        with zipfile.ZipFile(xlsx_path, "r") as zin, zipfile.ZipFile(tmp_name, "w", zipfile.ZIP_DEFLATED) as zout:
            for info in zin.infolist():
                data = zin.read(info.filename)
                if info.filename == sheet_xml_path:
                    xml = data.decode("utf-8")
                    if "</extLst>" in xml:
                        xml = xml.replace("</extLst>", ext_xml + "</extLst>", 1)
                    else:
                        xml = xml.replace("</worksheet>", "<extLst>" + ext_xml + "</extLst></worksheet>", 1)
                    data = xml.encode("utf-8")
                zout.writestr(info, data)
        os.replace(tmp_name, xlsx_path)
        return True
    finally:
        try:
            Path(tmp_name).unlink(missing_ok=True)
        except (FileNotFoundError, PermissionError):
            pass


def inject_x14_conditional_formatting(xlsx_path: Path, template_path: Path, sheet_targets: dict[str, list[dict]], layout, template_sheet_name=None):
    for sheet_name, targets in sheet_targets.items():
        inject_x14_conditional_formatting_for_sheet(xlsx_path, template_path, sheet_name, targets, layout, template_sheet_name)


# ===== drawing / image preservation for untouched sheets =====
def rels_path_for_part(part_path: str) -> str:
    part = PurePosixPath(part_path)
    return str(part.parent / "_rels" / f"{part.name}.rels")


def resolve_related_part(source_part: str, target: str) -> str | None:
    if not target or re.match(r"^[a-zA-Z][a-zA-Z0-9+.-]*:", target):
        return None
    if target.startswith("/"):
        return target.lstrip("/")
    return posixpath.normpath(str(PurePosixPath(source_part).parent / target)).lstrip("/")


def _read_zip_text(z, name: str) -> str | None:
    try:
        return z.read(name).decode("utf-8")
    except KeyError:
        return None


def _collect_internal_related_parts(z, part_path: str, collected: set[str]):
    if part_path in collected or part_path not in z.namelist():
        return
    collected.add(part_path)
    rels_path = rels_path_for_part(part_path)
    if rels_path not in z.namelist():
        return
    collected.add(rels_path)
    try:
        rels_root = ET.fromstring(z.read(rels_path))
    except Exception:
        return
    rel_ns = "{http://schemas.openxmlformats.org/package/2006/relationships}"
    for rel in rels_root.findall(f"{rel_ns}Relationship"):
        if rel.attrib.get("TargetMode") == "External":
            continue
        target_part = resolve_related_part(part_path, rel.attrib.get("Target", ""))
        if target_part:
            _collect_internal_related_parts(z, target_part, collected)


def _merge_content_types(src_xml: bytes, dst_xml: bytes, available_parts: set[str]) -> bytes:
    ns = "{http://schemas.openxmlformats.org/package/2006/content-types}"
    try:
        src_root = ET.fromstring(src_xml)
        dst_root = ET.fromstring(dst_xml)
    except Exception:
        return dst_xml

    existing_defaults = {el.attrib.get("Extension") for el in dst_root.findall(f"{ns}Default")}
    existing_overrides = {el.attrib.get("PartName") for el in dst_root.findall(f"{ns}Override")}
    for el in src_root.findall(f"{ns}Default"):
        if el.attrib.get("Extension") not in existing_defaults:
            dst_root.append(copy(el))
            existing_defaults.add(el.attrib.get("Extension"))
    for el in src_root.findall(f"{ns}Override"):
        if (el.attrib.get("PartName", "").lstrip("/") in available_parts
                and el.attrib.get("PartName") not in existing_overrides):
            dst_root.append(copy(el))
            existing_overrides.add(el.attrib.get("PartName"))
    ET.register_namespace("", "http://schemas.openxmlformats.org/package/2006/content-types")
    return ET.tostring(dst_root, encoding="utf-8", xml_declaration=True)


def _merge_workbook_rels_shared_strings(src_xml: bytes, dst_xml: bytes) -> bytes:
    rel_ns_uri = "http://schemas.openxmlformats.org/package/2006/relationships"
    rel_ns = f"{{{rel_ns_uri}}}"
    try:
        src_root = ET.fromstring(src_xml)
        dst_root = ET.fromstring(dst_xml)
    except Exception:
        return dst_xml

    shared_type_suffix = "/sharedStrings"
    src_shared = [
        rel for rel in src_root.findall(f"{rel_ns}Relationship")
        if str(rel.attrib.get("Type", "")).endswith(shared_type_suffix)
    ]
    if not src_shared:
        return dst_xml

    dst_shared = [
        rel for rel in dst_root.findall(f"{rel_ns}Relationship")
        if str(rel.attrib.get("Type", "")).endswith(shared_type_suffix)
    ]
    if dst_shared:
        return dst_xml

    used_ids = {rel.attrib.get("Id") for rel in dst_root.findall(f"{rel_ns}Relationship")}
    next_id = 1
    for src_rel in src_shared:
        attrs = dict(src_rel.attrib)
        rel_id = attrs.get("Id") or "rId1"
        if rel_id in used_ids:
            while f"rId{next_id}" in used_ids:
                next_id += 1
            rel_id = f"rId{next_id}"
        attrs["Id"] = rel_id
        dst_root.append(ET.Element(f"{rel_ns}Relationship", attrs))
        used_ids.add(rel_id)

    ET.register_namespace("", rel_ns_uri)
    return ET.tostring(dst_root, encoding="utf-8", xml_declaration=True)


def _sheet_xml_uses_shared_strings(xml: str) -> bool:
    return bool(re.search(r'\bt=["\']s["\']', xml or ""))


def collect_sheet_drawings_to_preserve(xlsx_path: Path, preserve_sheet_names: set[str]) -> dict[str, dict]:
    if not preserve_sheet_names or not xlsx_path.exists():
        return {}

    result = {}
    try:
        sheet_paths = sheet_xml_paths(xlsx_path)
    except Exception:
        return result

    drawing_tag_re = re.compile(
        r"<(?:\w+:)?(?:drawing|legacyDrawing|legacyDrawingHF|picture)\b[^>]*/>",
        re.IGNORECASE,
    )
    rid_re = re.compile(r'\br:id="([^"]+)"')
    rel_ns = "{http://schemas.openxmlformats.org/package/2006/relationships}"
    with zipfile.ZipFile(xlsx_path, "r") as z:
        zip_names = set(z.namelist())
        for sheet_name in preserve_sheet_names:
            xml_path = sheet_paths.get(sheet_name)
            if not xml_path:
                continue
            xml = _read_zip_text(z, xml_path)
            if not xml:
                continue
            drawing_tags = drawing_tag_re.findall(xml)
            if not drawing_tags:
                continue
            drawing_rids = set()
            for tag in drawing_tags:
                drawing_rids.update(rid_re.findall(tag))
            if not drawing_rids:
                continue

            rels_path = rels_path_for_part(xml_path)
            if rels_path not in zip_names:
                continue
            try:
                rels_root = ET.fromstring(z.read(rels_path))
            except Exception:
                continue
            rels = []
            copied_parts = set()
            for rel in rels_root.findall(f"{rel_ns}Relationship"):
                if rel.attrib.get("Id") not in drawing_rids:
                    continue
                rels.append(dict(rel.attrib))
                if rel.attrib.get("TargetMode") != "External":
                    target_part = resolve_related_part(xml_path, rel.attrib.get("Target", ""))
                    if target_part:
                        _collect_internal_related_parts(z, target_part, copied_parts)
            if rels:
                result[sheet_name] = {
                    "drawing_tags": drawing_tags,
                    "rels": rels,
                    "parts": copied_parts,
                }
    return result


def collect_sheet_parts_to_preserve(xlsx_path: Path, preserve_sheet_names: set[str]) -> dict[str, dict]:
    if not preserve_sheet_names or not xlsx_path.exists():
        return {}

    result = {}
    try:
        sheet_paths = sheet_xml_paths(xlsx_path)
    except Exception:
        return result

    with zipfile.ZipFile(xlsx_path, "r") as z:
        zip_names = set(z.namelist())
        for sheet_name in preserve_sheet_names:
            xml_path = sheet_paths.get(sheet_name)
            if not xml_path or xml_path not in zip_names:
                continue
            parts = set()
            _collect_internal_related_parts(z, xml_path, parts)
            parts.add(xml_path)
            xml = _read_zip_text(z, xml_path)
            if _sheet_xml_uses_shared_strings(xml or "") and "xl/sharedStrings.xml" in zip_names:
                parts.add("xl/sharedStrings.xml")
            result[sheet_name] = {"xml_path": xml_path, "parts": parts}
    return result


def sheet_names_with_image_parts(xlsx_path: Path) -> set[str]:
    if not xlsx_path.exists():
        return set()
    try:
        sheet_paths = sheet_xml_paths(xlsx_path)
    except Exception:
        return set()

    image_sheet_names = set()
    drawing_tag_re = re.compile(
        r"<(?:\w+:)?(?:drawing|legacyDrawing|legacyDrawingHF|picture|oleObjects?)\b",
        re.IGNORECASE,
    )
    rel_ns = "{http://schemas.openxmlformats.org/package/2006/relationships}"
    related_markers = (
        "/drawing",
        "/vmlDrawing",
        "/image",
        "/oleObject",
        "/control",
    )
    with zipfile.ZipFile(xlsx_path, "r") as z:
        zip_names = set(z.namelist())
        for sheet_name, xml_path in sheet_paths.items():
            xml = _read_zip_text(z, xml_path)
            if xml and drawing_tag_re.search(xml):
                image_sheet_names.add(sheet_name)
                continue
            rels_path = rels_path_for_part(xml_path)
            if rels_path not in zip_names:
                continue
            try:
                rels_root = ET.fromstring(z.read(rels_path))
            except Exception:
                continue
            for rel in rels_root.findall(f"{rel_ns}Relationship"):
                rel_type = rel.attrib.get("Type", "")
                target = rel.attrib.get("Target", "")
                if any(marker in rel_type or marker in target for marker in related_markers):
                    image_sheet_names.add(sheet_name)
                    break
    return image_sheet_names


def sheet_index_map(xlsx_path: Path) -> dict[str, int]:
    if not xlsx_path.exists():
        return {}
    try:
        with zipfile.ZipFile(xlsx_path, "r") as z:
            wb_root = ET.fromstring(z.read("xl/workbook.xml"))
    except Exception:
        return {}
    sheets = [el for el in wb_root.iter() if el.tag.endswith("sheet")]
    return {sheet.attrib.get("name"): idx for idx, sheet in enumerate(sheets) if sheet.attrib.get("name")}


def restore_sheet_order_positions(xlsx_path: Path, desired_indexes: dict[str, int]):
    if not desired_indexes:
        return False

    tmp_dir = xlsx_path.parent if xlsx_path.parent.exists() else Path(".")
    tmp = tempfile.NamedTemporaryFile(prefix=f".{xlsx_path.stem}_", suffix=".tmp.xlsx", dir=tmp_dir, delete=False)
    tmp_name = Path(tmp.name)
    tmp.close()
    changed = False
    ns_uri = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
    ET.register_namespace("", ns_uri)
    ET.register_namespace("r", "http://schemas.openxmlformats.org/officeDocument/2006/relationships")

    try:
        with zipfile.ZipFile(xlsx_path, "r") as zin, zipfile.ZipFile(tmp_name, "w", zipfile.ZIP_DEFLATED) as zout:
            for info in zin.infolist():
                data = zin.read(info.filename)
                if info.filename == "xl/workbook.xml":
                    try:
                        wb_root = ET.fromstring(data)
                        sheets_el = next((el for el in wb_root.iter() if el.tag.endswith("sheets")), None)
                    except Exception:
                        sheets_el = None
                    if sheets_el is not None:
                        sheet_elements = list(sheets_el)
                        preserved = [
                            sheet for sheet in sheet_elements
                            if sheet.attrib.get("name") in desired_indexes
                        ]
                        if preserved:
                            for sheet in preserved:
                                sheets_el.remove(sheet)
                            for sheet in sorted(preserved, key=lambda el: desired_indexes.get(el.attrib.get("name"), 10**9)):
                                insert_at = min(max(desired_indexes.get(sheet.attrib.get("name"), len(sheets_el)), 0), len(sheets_el))
                                sheets_el.insert(insert_at, sheet)
                            data = ET.tostring(wb_root, encoding="utf-8", xml_declaration=True)
                            changed = True
                zout.writestr(info, data)
        if changed:
            os.replace(tmp_name, xlsx_path)
        return changed
    finally:
        try:
            tmp_name.unlink(missing_ok=True)
        except Exception:
            pass


def reinject_preserved_sheet_parts(xlsx_path: Path, source_path: Path, preserved: dict[str, dict]):
    if not preserved:
        return False
    try:
        output_sheet_paths = sheet_xml_paths(xlsx_path)
        source_sheet_paths = sheet_xml_paths(source_path)
    except Exception:
        return False

    source_to_output = {}
    parts_to_copy = set()
    for sheet_name, item in preserved.items():
        src_xml_path = source_sheet_paths.get(sheet_name, item.get("xml_path"))
        dst_xml_path = output_sheet_paths.get(sheet_name)
        if not src_xml_path or not dst_xml_path:
            continue
        source_to_output[src_xml_path] = dst_xml_path
        for part in item.get("parts", set()):
            if part == src_xml_path:
                parts_to_copy.add((src_xml_path, dst_xml_path))
            else:
                parts_to_copy.add((part, part))
        src_rels = rels_path_for_part(src_xml_path)
        dst_rels = rels_path_for_part(dst_xml_path)
        parts_to_copy.add((src_rels, dst_rels))

    if not parts_to_copy:
        return False

    tmp_dir = xlsx_path.parent if xlsx_path.parent.exists() else Path(".")
    tmp = tempfile.NamedTemporaryFile(prefix=f".{xlsx_path.stem}_", suffix=".tmp.xlsx", dir=tmp_dir, delete=False)
    tmp_name = Path(tmp.name)
    tmp.close()

    try:
        with zipfile.ZipFile(source_path, "r") as zsrc, zipfile.ZipFile(xlsx_path, "r") as zin, zipfile.ZipFile(tmp_name, "w", zipfile.ZIP_DEFLATED) as zout:
            source_names = set(zsrc.namelist())
            replacement_by_dest = {dst: src for src, dst in parts_to_copy if src in source_names}
            shared_strings_needed = "xl/sharedStrings.xml" in replacement_by_dest
            copied_dest = set()
            for info in zin.infolist():
                if info.filename == "[Content_Types].xml":
                    data = _merge_content_types(
                        zsrc.read(info.filename), zin.read(info.filename),
                        set(zin.namelist()) | set(replacement_by_dest),
                    )
                elif info.filename == "xl/_rels/workbook.xml.rels" and shared_strings_needed and info.filename in source_names:
                    data = _merge_workbook_rels_shared_strings(zsrc.read(info.filename), zin.read(info.filename))
                elif info.filename in replacement_by_dest:
                    data = zsrc.read(replacement_by_dest[info.filename])
                    copied_dest.add(info.filename)
                else:
                    data = zin.read(info.filename)
                zout.writestr(info, data)

            for dest, src in replacement_by_dest.items():
                if dest in copied_dest:
                    continue
                zout.writestr(dest, zsrc.read(src))
                copied_dest.add(dest)

        os.replace(tmp_name, xlsx_path)
        return True
    finally:
        try:
            tmp_name.unlink(missing_ok=True)
        except Exception:
            pass


def reinject_preserved_sheet_drawings(xlsx_path: Path, source_path: Path, preserved: dict[str, dict]):
    if not preserved:
        return False
    try:
        sheet_paths = sheet_xml_paths(xlsx_path)
    except Exception:
        return False

    tmp_dir = xlsx_path.parent if xlsx_path.parent.exists() else Path(".")
    tmp = tempfile.NamedTemporaryFile(prefix=f".{xlsx_path.stem}_", suffix=".tmp.xlsx", dir=tmp_dir, delete=False)
    tmp_name = Path(tmp.name)
    tmp.close()
    rel_ns_uri = "http://schemas.openxmlformats.org/package/2006/relationships"
    rel_ns = f"{{{rel_ns_uri}}}"
    ET.register_namespace("", rel_ns_uri)
    changed = False
    copied_parts = set()

    try:
        with zipfile.ZipFile(source_path, "r") as zsrc, zipfile.ZipFile(xlsx_path, "r") as zin, zipfile.ZipFile(tmp_name, "w", zipfile.ZIP_DEFLATED) as zout:
            source_names = set(zsrc.namelist())
            target_names = set(zin.namelist())
            xml_replacements = {}
            rel_replacements = {}

            for sheet_name, item in preserved.items():
                xml_path = sheet_paths.get(sheet_name)
                if not xml_path or xml_path not in target_names:
                    continue
                xml = zin.read(xml_path).decode("utf-8")
                xml = re.sub(r"<(?:\w+:)?(?:drawing|legacyDrawing|legacyDrawingHF|picture)\b[^>]*/>", "", xml)

                rels_path = rels_path_for_part(xml_path)
                if rels_path in target_names:
                    rels_root = ET.fromstring(zin.read(rels_path))
                else:
                    rels_root = ET.Element(f"{rel_ns}Relationships")

                used_ids = {rel.attrib.get("Id") for rel in rels_root.findall(f"{rel_ns}Relationship")}
                next_id = 1
                tag_xml = "\n".join(item["drawing_tags"])
                for rel_attrs in item["rels"]:
                    old_id = rel_attrs.get("Id")
                    new_id = old_id
                    if new_id in used_ids:
                        while f"rId{next_id}" in used_ids:
                            next_id += 1
                        new_id = f"rId{next_id}"
                    attrs = dict(rel_attrs)
                    attrs["Id"] = new_id
                    ET.SubElement(rels_root, f"{rel_ns}Relationship", attrs)
                    used_ids.add(new_id)
                    if old_id and old_id != new_id:
                        tag_xml = tag_xml.replace(f'r:id="{old_id}"', f'r:id="{new_id}"')

                insert_at = xml.find("<extLst")
                if insert_at < 0:
                    insert_at = xml.rfind("</worksheet>")
                if insert_at >= 0:
                    xml = xml[:insert_at] + tag_xml + xml[insert_at:]
                    xml_replacements[xml_path] = xml.encode("utf-8")
                    rel_replacements[rels_path] = ET.tostring(rels_root, encoding="utf-8", xml_declaration=True)
                    copied_parts.update(item["parts"])
                    changed = True

            for info in zin.infolist():
                if info.filename == "[Content_Types].xml" and copied_parts:
                    data = _merge_content_types(
                        zsrc.read(info.filename), zin.read(info.filename),
                        target_names | (copied_parts & source_names),
                    )
                elif info.filename in xml_replacements:
                    data = xml_replacements[info.filename]
                elif info.filename in rel_replacements:
                    data = rel_replacements[info.filename]
                else:
                    data = zin.read(info.filename)
                zout.writestr(info, data)

            existing_after_base = target_names | set(xml_replacements) | set(rel_replacements)
            for part in sorted(copied_parts):
                if part in existing_after_base or part not in source_names:
                    continue
                zout.writestr(part, zsrc.read(part))
                existing_after_base.add(part)
            for rels_path, data in rel_replacements.items():
                if rels_path not in target_names:
                    zout.writestr(rels_path, data)

        if changed:
            os.replace(tmp_name, xlsx_path)
        return changed
    finally:
        try:
            tmp_name.unlink(missing_ok=True)
        except Exception:
            pass


# ===== summary =====
def quote_sheet_name(sheet_name):
    return "'" + str(sheet_name).replace("'", "''") + "'"


def _normalize_compare_key(value):
    text = "" if value is None else str(value)
    return text.strip().lower()


def _normalize_table_key(value):
    text = _normalize_compare_key(value)
    if text.endswith(".csv"):
        text = text[:-4]
    return text


def _normalize_mark_text(value):
    text = "" if value is None else str(value).strip()
    return "CIRCLE" if text in {"\u25cb", "\u3007"} else text


def is_selected_mark(value, expected):
    return _normalize_mark_text(value) == _normalize_mark_text(expected)


def read_compare_exclusions(wb):
    exclusions = set()
    if COMPARE_EXCLUDE_SHEET_NAME not in wb.sheetnames:
        return exclusions

    ws = wb[COMPARE_EXCLUDE_SHEET_NAME]
    table_col = 1
    item_col = 2
    for c in range(1, min(ws.max_column, 20) + 1):
        header = _normalize_compare_key(ws.cell(1, c).value)
        if header in {"テーブル名", "table", "table_name", "表名"}:
            table_col = c
        elif header in {"項目名", "项目名", "column", "column_name", "field", "field_name"}:
            item_col = c

    for r in range(2, ws.max_row + 1):
        table_name = ws.cell(r, table_col).value
        item_name = ws.cell(r, item_col).value
        item_key = _normalize_compare_key(item_name)
        if not item_key:
            continue
        table_key = _normalize_table_key(table_name)
        if not table_key:
            table_key = "*"
        exclusions.add((table_key, item_key))
    return exclusions


def make_table_aliases(*values):
    aliases = set()
    for value in values:
        if value is None:
            continue
        text = str(value).strip()
        if not text:
            continue
        aliases.add(_normalize_table_key(text))
        aliases.add(_normalize_compare_key(text))
        if text.lower().endswith(".csv"):
            aliases.add(_normalize_table_key(text[:-4]))
    return aliases


def is_compare_excluded(table_aliases, item_name, compare_exclusions):
    item_key = _normalize_compare_key(item_name)
    if not item_key or not compare_exclusions:
        return False
    if ("*", item_key) in compare_exclusions:
        return True
    aliases = table_aliases if isinstance(table_aliases, set) else make_table_aliases(*(table_aliases or []))
    return any((alias, item_key) in compare_exclusions for alias in aliases)


def xl_abs_cell(sheet_name, row, col):
    return f"{quote_sheet_name(sheet_name)}!${get_column_letter(col)}${row}"


def xl_abs_range(sheet_name, row1, col1, row2, col2):
    return f"{quote_sheet_name(sheet_name)}!${get_column_letter(col1)}${row1}:${get_column_letter(col2)}${row2}"


def build_non_empty_row_count_formula(sheet_name, row1, row2, col1, col2):
    data_range = xl_abs_range(sheet_name, row1, col1, row2, col2)
    first_row_range = xl_abs_range(sheet_name, row1, col1, row1, col2)
    return ("=IFERROR("
            f"SUMPRODUCT(--(MMULT(--({data_range}<>\"\"),TRANSPOSE(COLUMN({first_row_range})^0))>0)),"
            "0)")


def build_countif_diff_formula(ranges):
    if not ranges:
        return '="OK"'
    parts = [f'COUNTIF({rng},"差分あり")' for rng in ranges]
    return f'=IF(({"+".join(parts)})>0,"NG","OK")'


def write_summary_sheet(wb, records, sheet_index=0):
    for old_name in (SUMMARY_SHEET_NAME, "summary"):
        if old_name in wb.sheetnames:
            del wb[old_name]
    safe_index = min(max(0, sheet_index or 0), len(wb.worksheets))
    ws = wb.create_sheet(SUMMARY_SHEET_NAME, safe_index)
    headers = [
        "CRUD", "出力番号", "出力Sheet", "テーブル名",
        "現_打鍵前件数", "現_打鍵後件数", "現_打鍵前後件数差異",
        "新_打鍵前件数", "新_打鍵後件数", "新_打鍵前後件数差異",
        "現_打鍵前CSVパス", "現_打鍵後CSVパス",
        "現新件数チェック", "差分チェック", "備考",
    ]
    ws.append(headers)
    for c in range(1, len(headers) + 1):
        cell = ws.cell(1, c)
        cell.fill = copy(SUMMARY_HEADER_FILL)
        cell.font = Font(name=FONT_NAME, bold=True, color="FFFFFF")
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = Border(left=THIN_SIDE, right=THIN_SIDE, top=THIN_SIDE, bottom=THIN_SIDE)

    for r_idx, rec in enumerate(records, start=2):
        refs = rec.detail_refs or {}
        sheet_name = rec.output_sheet
        if refs:
            gen_before_formula = f"=IFERROR({xl_abs_cell(sheet_name, refs['before_count_row'], 3)},0)"
            gen_after_formula = f"=IFERROR({xl_abs_cell(sheet_name, refs['after_count_row'], 3)},0)"
            new_before_formula = build_non_empty_row_count_formula(sheet_name, refs["before_new_start"], refs["before_new_end"], DATA_START_COL, refs["data_last_col"])
            new_after_formula = build_non_empty_row_count_formula(sheet_name, refs["after_new_start"], refs["after_new_end"], DATA_START_COL, refs["data_last_col"])
            diff_ranges = [xl_abs_range(sheet_name, s, 2, e, 2) for s, e in refs.get("compare_b_ranges", []) if s and e and e >= s]
            diff_check_formula = build_countif_diff_formula(diff_ranges)
        else:
            gen_before_formula = rec.before.count
            gen_after_formula = rec.after.count
            if rec.current_no_update:
                new_before_formula = rec.before.count
                new_after_formula = rec.after.count
            else:
                new_before_formula = 0
                new_after_formula = 0
            diff_check_formula = '="OK"'

        values = [
            "" if rec.current_no_update else f'=IF(F{r_idx}>E{r_idx},"C",IF(F{r_idx}<E{r_idx},"D",IF(AND(F{r_idx}=E{r_idx},E{r_idx}<>0),"U","")))',
            rec.output_no if rec.output_no is not None else "",
            rec.output_sheet,
            rec.display_name,
            gen_before_formula,
            gen_after_formula,
            f"=F{r_idx}-E{r_idx}",
            new_before_formula,
            new_after_formula,
            f"=I{r_idx}-H{r_idx}",
            str(rec.before.path) if rec.before.path else "",
            str(rec.after.path) if rec.after.path else "",
            f'=IF(AND(E{r_idx}=H{r_idx},F{r_idx}=I{r_idx}),"差分なし","差分あり")',
            diff_check_formula,
            " / ".join(rec.notes),
        ]
        ws.append(values)
        row_fill = {"C": SUMMARY_C_FILL, "U": SUMMARY_U_FILL, "D": SUMMARY_D_FILL, "": SUMMARY_SKIP_FILL}.get(rec.crud, SUMMARY_SKIP_FILL)
        for c in range(1, len(headers) + 1):
            cell = ws.cell(r_idx, c)
            cell.fill = copy(row_fill)
            cell.font = Font(name=FONT_NAME)
            cell.alignment = Alignment(vertical="center", wrap_text=(c in {11, 12, len(headers)}))
            cell.border = Border(left=THIN_SIDE, right=THIN_SIDE, top=THIN_SIDE, bottom=THIN_SIDE)
        if rec.crud == "":
            ws.row_dimensions[r_idx].hidden = True

    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{max(1, ws.max_row)}"
    ws.freeze_panes = "A2"
    widths = [10, 12, 16, 32, 18, 18, 22, 18, 18, 22, 45, 45, 18, 16, 70]
    for i, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = width
    if ws.max_row >= 2:
        count_check_ref = f"M2:M{ws.max_row}"
        data_check_ref = f"N2:N{ws.max_row}"
        ws.conditional_formatting.add(count_check_ref, CellIsRule(operator="equal", formula=['"差分なし"'], fill=TRUE_FILL))
        ws.conditional_formatting.add(count_check_ref, CellIsRule(operator="equal", formula=['"差分あり"'], fill=FALSE_FILL, font=FALSE_FONT))
        ws.conditional_formatting.add(data_check_ref, CellIsRule(operator="equal", formula=['"OK"'], fill=TRUE_FILL))
        ws.conditional_formatting.add(data_check_ref, CellIsRule(operator="equal", formula=['"NG"'], fill=FALSE_FILL, font=FALSE_FONT))
    for r in range(1, ws.max_row + 1):
        ws.row_dimensions[r].height = HEADER_ROW_HEIGHT if r == 1 else BODY_ROW_HEIGHT
    return ws


# ===== workbook =====
def normalize_fonts_and_row_heights(ws, max_row, max_col, header_rows):
    for r in range(1, max_row + 1):
        ws.row_dimensions[r].height = HEADER_ROW_HEIGHT if r in header_rows else BODY_ROW_HEIGHT
        for c in range(1, max_col + 1):
            cell = ws.cell(r, c)
            if not isinstance(cell, MergedCell):
                force_font_name(cell)


def copy_sheet_page_settings(tws, ows):
    try:
        ows.sheet_view.showGridLines = tws.sheet_view.showGridLines
    except Exception:
        pass
    try:
        ows.freeze_panes = tws.freeze_panes
    except Exception:
        pass
    for attr in ("page_margins", "page_setup", "print_options", "sheet_properties"):
        try:
            setattr(ows, attr, copy(getattr(tws, attr)))
        except Exception:
            pass


def unique_sheet_name(wb, base_name):
    if base_name not in wb.sheetnames:
        return base_name
    i = 1
    while f"{base_name}{i}" in wb.sheetnames:
        i += 1
    return f"{base_name}{i}"


def assign_output_positions(records, base_sheet_name, tables_per_sheet):
    output_records = [r for r in records if not r.skip_detail]
    multi = len(output_records) > tables_per_sheet
    for idx, rec in enumerate(output_records, start=1):
        rec.output_no = idx
        rec.output_sheet = f"{base_sheet_name}{(idx - 1) // tables_per_sheet + 1}" if multi else base_sheet_name
    return output_records



# ===== 設定ファイル / 出力制御 =====
DEFAULT_CONFIG_FILE = "設定ファイル.json"
ACTION_REPLACE_FILE = "replace_file"
ACTION_ERROR = "error"
ACTION_REPLACE_SHEETS = "replace_sheets"


def setup_logging(script_name: str):
    log_dir = Path.cwd() / "log"
    log_dir.mkdir(parents=True, exist_ok=True)
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    log_path = log_dir / f"{script_name}_{timestamp}.log"
    logging.basicConfig(
        level=logging.INFO,
        format="%(asctime)s [%(levelname)s] %(message)s",
        handlers=[logging.FileHandler(log_path, encoding="utf-8")],
        force=True,
    )
    return logging.getLogger(script_name), log_path


def print_and_log(message: str, logger=None):
    print(message)
    if logger:
        logger.info(message)


def load_settings_file(config_path: Path | None):
    if config_path is None:
        config_path = Path(DEFAULT_CONFIG_FILE)
    if not config_path.is_absolute():
        local = Path.cwd() / config_path
        script_local = Path(__file__).resolve().parent / config_path
        config_path = local if local.exists() else script_local
    if not config_path.exists():
        return {}
    with config_path.open("r", encoding="utf-8-sig") as f:
        return json.load(f)


def cfg_get(*dicts, keys, default=None):
    for d in dicts:
        if not isinstance(d, dict):
            continue
        for key in keys:
            if key in d and d[key] not in (None, ""):
                return d[key]
    return default


def bool_value(value, default=False):
    if value is None:
        return default
    if isinstance(value, bool):
        return value
    return str(value).strip().lower() in {"1", "true", "yes", "y", "on", "有", "あり", "する", "はい"}


def normalize_existing_action(value):
    text = str(value or ACTION_REPLACE_FILE).strip().lower()
    mapping = {
        "replace": ACTION_REPLACE_FILE,
        "overwrite": ACTION_REPLACE_FILE,
        "replace_file": ACTION_REPLACE_FILE,
        "上書き": ACTION_REPLACE_FILE,
        "ファイル上書き": ACTION_REPLACE_FILE,
        "error": ACTION_ERROR,
        "raise": ACTION_ERROR,
        "エラー": ACTION_ERROR,
        "replace_sheets": ACTION_REPLACE_SHEETS,
        "sheet": ACTION_REPLACE_SHEETS,
        "sheets": ACTION_REPLACE_SHEETS,
        "シート差替": ACTION_REPLACE_SHEETS,
        "シート差し替え": ACTION_REPLACE_SHEETS,
    }
    if text not in mapping:
        raise ValueError(f"既存ファイル処理が不正です: {value}（replace_file / error / replace_sheets）")
    return mapping[text]


def normalize_path_value(value):
    """設定ファイル上のパス表記を正規化する。

    JSON では通常どおり "." / ".." を使えるほか、説明用に
    「現在フォルダ」「上一层文件夹」などを書いた場合も解釈する。
    """
    if value is None:
        return "."
    text = str(value).strip()
    aliases = {
        "": ".",
        "現在フォルダ": ".",
        "現在フォルダー": ".",
        "カレントフォルダ": ".",
        "カレントフォルダー": ".",
        "current": ".",
        "current_folder": ".",
        "当前文件夹": ".",
        "当前目录": ".",
        "上一层文件夹": "..",
        "上一層文件夾": "..",
        "上位フォルダ": "..",
        "上位フォルダー": "..",
        "親フォルダ": "..",
        "親フォルダー": "..",
        "parent": "..",
        "parent_folder": "..",
    }
    return aliases.get(text, text)


def join_if_relative(base, value):
    value = normalize_path_value(value)
    base = normalize_path_value(base)
    path = Path(str(value))
    if path.is_absolute():
        return path
    return Path(str(base or ".")) / path



def _as_ordered_values(value):
    """設定値を優先順のリストへ変換する。

    JSON 配列のほか、文字列では ; / ； / 改行区切りも利用できる。
    Windows の C:/... を壊さないため、コロンでは分割しない。
    """
    if value is None:
        return ["."]
    if isinstance(value, (list, tuple)):
        result = []
        for item in value:
            result.extend(_as_ordered_values(item))
        return result or ["."]
    text = str(value).strip()
    if text == "":
        return ["."]
    parts = re.split(r"[;；\r\n]+", text)
    parts = [p.strip() for p in parts if p and p.strip()]
    return parts or [text]


def _first_ordered_value(value, default=None):
    values = _as_ordered_values(value)
    if not values:
        return default
    return values[0]


def has_any_key(*dicts, keys):
    for d in dicts:
        if not isinstance(d, dict):
            continue
        for key in keys:
            if key in d and d[key] not in (None, ""):
                return True
    return False


def _dedupe_paths(paths):
    result = []
    seen = set()
    for p in paths:
        p = Path(str(normalize_path_value(p)))
        try:
            key = str(p.resolve()) if p.exists() else str(p.absolute())
        except Exception:
            key = str(p)
        if key in seen:
            continue
        seen.add(key)
        result.append(p)
    return result


def resolve_path_list(path_value):
    """設定ファイル上のCSVフォルダパスを優先順リストとして解決する。

    例:
      "現_打鍵前CSVパス": ["優先1/現_打鍵前", "優先2/現_打鍵前"]
      => 優先1/現_打鍵前, 優先2/現_打鍵前 の順で検索する。
    """
    return _dedupe_paths(_as_ordered_values(path_value))


def resolve_folder_list(base_value, folder_value):
    """旧設定互換: 入力パス + フォルダ名を優先順リストとして解決する。"""
    bases = _as_ordered_values(base_value)
    folders = _as_ordered_values(folder_value)
    return _dedupe_paths(join_if_relative(base, folder) for base in bases for folder in folders)


def resolve_file_path(path_value, must_exist=False):
    """設定ファイル上のファイルパスを解決する。ワイルドカードは1件一致のみ許可する。"""
    raw = _first_ordered_value(path_value, default=".")
    file_path = Path(str(normalize_path_value(raw)))
    pattern = str(file_path)
    if any(ch in pattern for ch in "*?["):
        parent = file_path.parent if str(file_path.parent) else Path(".")
        matches = sorted(parent.glob(file_path.name))
        matches = [m for m in matches if m.is_file()]
        if len(matches) == 1:
            return matches[0]
        if not matches:
            raise FileNotFoundError(f"ワイルドカードに一致するファイルがありません: {pattern}")
        raise RuntimeError("ワイルドカードに一致するファイルが複数あります: " + ", ".join(str(m) for m in matches))
    if must_exist and not file_path.exists():
        raise FileNotFoundError(f"ファイルが存在しません: {file_path}")
    return file_path

def resolve_wildcard_path(folder, filename, must_exist=False):
    base = Path(str(normalize_path_value(folder or ".")))
    pattern_path = join_if_relative(base, filename)
    pattern = str(pattern_path)
    if any(ch in pattern for ch in "*?["):
        matches = sorted(Path(pattern_path.parent).glob(pattern_path.name))
        matches = [m for m in matches if m.is_file()]
        if len(matches) == 1:
            return matches[0]
        if not matches:
            raise FileNotFoundError(f"ワイルドカードに一致するファイルがありません: {pattern}")
        raise RuntimeError("ワイルドカードに一致するファイルが複数あります: " + ", ".join(str(m) for m in matches))
    if must_exist and not pattern_path.exists():
        raise FileNotFoundError(f"ファイルが存在しません: {pattern_path}")
    return pattern_path

def column_index_from_config(value, label):
    """設定値の列指定を 1 始まりの列番号へ変換する。例: K / k / K列 / 11。"""
    if value is None or str(value).strip() == "":
        raise ValueError(f"{label} が設定されていません。")
    text = str(value).strip()
    if text.endswith("列"):
        text = text[:-1].strip()
    if text.isdigit():
        idx = int(text)
        if idx <= 0:
            raise ValueError(f"{label} は 1 以上を指定してください: {value}")
        return idx
    text = text.upper()
    if not re.fullmatch(r"[A-Z]+", text):
        raise ValueError(f"{label} の列指定が不正です: {value}")
    idx = 0
    for ch in text:
        idx = idx * 26 + (ord(ch) - ord("A") + 1)
    return idx


def read_selected_table_names_from_workbook(path: Path, sheet_name, mark_col, table_col, mark_value="〇"):
    """出力済みExcelの指定シートから、〇が付いた表名を読む。

    指定シートが存在しない場合は None を返し、呼び出し側はフィルタなしとして扱う。
    """
    if not sheet_name or path is None or not path.exists():
        return None
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        if sheet_name not in wb.sheetnames:
            return None
        ws = wb[sheet_name]
        mark_col_idx = column_index_from_config(mark_col, "対象表選択判定列")
        table_col_idx = column_index_from_config(table_col, "対象表選択表名列")
        expected = "" if mark_value is None else str(mark_value).strip()
        selected = set()
        for r in range(1, ws.max_row + 1):
            mark = ws.cell(r, mark_col_idx).value
            table_name = ws.cell(r, table_col_idx).value
            if not is_selected_mark(mark, expected):
                continue
            key = _normalize_table_key(table_name)
            if key:
                selected.add(key)
        return selected
    finally:
        try:
            wb.close()
        except Exception:
            pass


def read_selected_table_conditions_from_workbook(path: Path, sheet_name, mark_col, table_col, before_current_col, before_new_col, after_current_col, after_new_col, mark_value="〇"):
    """出力済みExcelの指定シートから、〇が付いた表名ごとの4種類の抽出条件を読む。"""
    condition_cols = [before_current_col, before_new_col, after_current_col, after_new_col]
    if not sheet_name or not any(condition_cols) or path is None or not path.exists():
        return {}
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        if sheet_name not in wb.sheetnames:
            return {}
        ws = wb[sheet_name]
        mark_col_idx = column_index_from_config(mark_col, "対象表選択判定列")
        table_col_idx = column_index_from_config(table_col, "対象表選択表名列")
        before_current_idx = column_index_from_config(before_current_col, "対象表選択打鍵前現条件列") if before_current_col else None
        before_new_idx = column_index_from_config(before_new_col, "対象表選択打鍵前新条件列") if before_new_col else None
        after_current_idx = column_index_from_config(after_current_col, "対象表選択打鍵後現条件列") if after_current_col else None
        after_new_idx = column_index_from_config(after_new_col, "対象表選択打鍵後新条件列") if after_new_col else None
        expected = "" if mark_value is None else str(mark_value).strip()
        conditions = {}

        def read_condition(row, col_idx):
            if not col_idx:
                return ""
            value = ws.cell(row, col_idx).value
            return "" if value is None else str(value).strip()

        for r in range(1, ws.max_row + 1):
            mark = ws.cell(r, mark_col_idx).value
            table_name = ws.cell(r, table_col_idx).value
            if not is_selected_mark(mark, expected):
                continue
            key = _normalize_table_key(table_name)
            if not key:
                continue
            conditions[key] = {
                "before_current": read_condition(r, before_current_idx),
                "before_new": read_condition(r, before_new_idx),
                "after_current": read_condition(r, after_current_idx),
                "after_new": read_condition(r, after_new_idx),
            }
        return conditions
    finally:
        try:
            wb.close()
        except Exception:
            pass


def record_matches_selected_tables(record, selected_table_names):
    if selected_table_names is None:
        return True
    aliases = make_table_aliases(record.display_name, record.sql_name, record.key)
    return any(alias in selected_table_names for alias in aliases)


def apply_selected_table_conditions(records, table_conditions):
    if not table_conditions:
        return
    for record in records:
        aliases = make_table_aliases(record.display_name, record.sql_name, record.key)
        for alias in aliases:
            if alias in table_conditions:
                conditions = table_conditions[alias]
                if isinstance(conditions, dict):
                    record.before_current_extraction_condition = conditions.get("before_current", "")
                    record.before_new_extraction_condition = conditions.get("before_new", "")
                    record.after_current_extraction_condition = conditions.get("after_current", "")
                    record.after_new_extraction_condition = conditions.get("after_new", "")
                else:
                    record.before_current_extraction_condition = str(conditions or "")
                    record.before_new_extraction_condition = str(conditions or "")
                    record.after_current_extraction_condition = str(conditions or "")
                    record.after_new_extraction_condition = str(conditions or "")
                break


def table_aliases_with_csv_data_from_records(records):
    aliases = set()
    for record in records:
        if record.before.count <= 0 and record.after.count <= 0:
            continue
        aliases.update(make_table_aliases(record.display_name, record.sql_name, record.key))
    return aliases


def find_or_create_selector_result_col(ws, header_name="CSVデータ不一致", header_row=3):
    for c in range(1, ws.max_column + 1):
        if str(ws.cell(header_row, c).value or "").strip() == header_name:
            return c
    for c in range(1, ws.max_column + 1):
        if str(ws.cell(1, c).value or "").strip() == header_name:
            ws.cell(1, c).value = None
            ws.cell(header_row, c).value = header_name
            ws.cell(header_row, c).font = Font(name=FONT_NAME, bold=True)
            ws.cell(header_row, c).alignment = Alignment(horizontal="center", vertical="center")
            return c
    col = ws.max_column + 1
    ws.cell(header_row, col).value = header_name
    ws.cell(header_row, col).font = Font(name=FONT_NAME, bold=True)
    ws.cell(header_row, col).alignment = Alignment(horizontal="center", vertical="center")
    ws.column_dimensions[get_column_letter(col)].width = max(ws.column_dimensions[get_column_letter(col)].width or 0, 28)
    return col


def find_or_create_selector_note_col(ws, result_col_idx, header_name="備考", header_row=3):
    target_col = result_col_idx + 1
    if str(ws.cell(header_row, target_col).value or "").strip() == header_name:
        return target_col
    if not str(ws.cell(header_row, target_col).value or "").strip():
        note_col_idx = target_col
    else:
        ws.insert_cols(target_col)
        note_col_idx = target_col
    ws.cell(header_row, note_col_idx).value = header_name
    ws.cell(header_row, note_col_idx).font = Font(name=FONT_NAME, bold=True)
    ws.cell(header_row, note_col_idx).alignment = Alignment(horizontal="center", vertical="center")
    ws.column_dimensions[get_column_letter(note_col_idx)].width = max(ws.column_dimensions[get_column_letter(note_col_idx)].width or 0, 36)
    return note_col_idx


def table_aliases_with_current_no_update(records):
    aliases = set()
    for record in records:
        if record.current_no_update:
            aliases.update(make_table_aliases(record.display_name, record.sql_name, record.key))
    return aliases


def selector_notes_by_table_alias(records):
    notes_by_alias = {}
    for record in records:
        selector_notes = []
        if record.current_no_update:
            selector_notes.append(CURRENT_NO_UPDATE_NOTE)
        selector_notes.extend(f"打鍵前CSVデータ補正: {note}" for note in record.before.repair_notes)
        selector_notes.extend(f"打鍵後CSVデータ補正: {note}" for note in record.after.repair_notes)
        if not selector_notes:
            continue
        for alias in make_table_aliases(record.display_name, record.sql_name, record.key):
            notes_by_alias[alias] = " / ".join(dict.fromkeys(selector_notes))
    return notes_by_alias


def mark_selector_csv_mismatches(wb, sheet_name, mark_col, table_col, mark_value, data_table_aliases, selected_table_names=None, selector_note_by_alias=None):
    if not sheet_name or sheet_name not in wb.sheetnames:
        return 0
    ws = wb[sheet_name]
    mark_col_idx = column_index_from_config(mark_col, "対象表選択判定列")
    table_col_idx = column_index_from_config(table_col, "対象表選択表名列")
    result_col_idx = find_or_create_selector_result_col(ws, header_row=3)
    note_col_idx = find_or_create_selector_note_col(ws, result_col_idx, header_row=3)
    expected = "" if mark_value is None else str(mark_value).strip()
    selector_note_by_alias = selector_note_by_alias or {}
    mismatch_count = 0
    for r in range(4, ws.max_row + 1):
        table_name = ws.cell(r, table_col_idx).value
        table_key = _normalize_table_key(table_name)
        if not table_key:
            ws.cell(r, result_col_idx).value = None
            ws.cell(r, note_col_idx).value = None
            continue
        if selected_table_names is not None:
            selected = table_key in selected_table_names
        else:
            selected = is_selected_mark(ws.cell(r, mark_col_idx).value, expected)
        has_csv_data = table_key in data_table_aliases
        result = ""
        if selected and not has_csv_data:
            result = "選択あり/CSVデータなし"
        elif not selected and has_csv_data:
            result = "CSVデータあり/選択なし"
        ws.cell(r, result_col_idx).value = result
        force_font_name(ws.cell(r, result_col_idx))
        note = selector_note_by_alias.get(table_key, "")
        ws.cell(r, note_col_idx).value = note
        force_font_name(ws.cell(r, note_col_idx))
        if result:
            mismatch_count += 1
    return mismatch_count

def east_asian_len(value):
    import unicodedata
    text = "" if value is None else str(value)
    total = 0
    for ch in text:
        width = unicodedata.east_asian_width(ch)
        total += 2 if width in {"F", "W"} else 1
    return total


def _is_sql_like_text(text):
    upper = text.upper()
    return "SELECT" in upper and (" FROM " in upper or " WHERE " in upper or "抽出条件" in text)


def _should_skip_for_autofit(ws, cell):
    value = cell.value
    if value in (None, ""):
        return True
    if isinstance(value, str):
        text = value.strip()
        if not text:
            return True
        # 数式や取得SQLは列幅を極端に広げる原因になるため、列幅計算から除外する。
        if text.startswith("=") or _is_sql_like_text(text):
            return True
    return False


def _effective_autofit_len(value):
    text = "" if value is None else str(value).replace("\r", "\n")
    # 改行がある場合は最長行だけを見る。備考などの長文が列幅を押し広げすぎないようにする。
    parts = text.split("\n")
    return max((east_asian_len(part) for part in parts), default=0)


def auto_fit_column_widths(wb, enabled=False, max_width=60, min_width=8, sheet_names=None):
    if not enabled:
        return
    try:
        max_width = float(max_width)
        min_width = float(min_width)
    except Exception:
        max_width, min_width = 60, 8
    if max_width < min_width:
        max_width, min_width = min_width, max_width

    target_names = set(sheet_names or wb.sheetnames)
    for ws in wb.worksheets:
        if ws.title not in target_names:
            continue
        for col in range(1, ws.max_column + 1):
            max_len = 0
            for row in range(1, ws.max_row + 1):
                if getattr(ws.row_dimensions[row], "hidden", False):
                    continue
                cell = ws.cell(row, col)
                if isinstance(cell, MergedCell) or _should_skip_for_autofit(ws, cell):
                    continue
                max_len = max(max_len, _effective_autofit_len(cell.value))

            if max_len:
                # Excel の列幅は実文字数より少し余裕を持たせるだけにする。
                # 以前の max_len+2 は SQL 行や長い式に引っ張られて 40 超になりやすかったため抑制する。
                width = max_len * 1.08 + 1.5
                ws.column_dimensions[get_column_letter(col)].width = min(max_width, max(min_width, width))

def is_table_output_sheet(name, prefix):
    return name == prefix or re.fullmatch(re.escape(prefix) + r"\d+", name or "") is not None


def remove_existing_output_sheets(wb, table_prefix):
    indexes = {ws.title: idx for idx, ws in enumerate(wb.worksheets)}
    table_indexes = [idx for idx, ws in enumerate(wb.worksheets) if is_table_output_sheet(ws.title, table_prefix)]
    summary_index = indexes.get(SUMMARY_SHEET_NAME, indexes.get("summary", 0))
    table_index = min(table_indexes) if table_indexes else len(wb.worksheets)
    delete_titles = [ws.title for ws in wb.worksheets if is_table_output_sheet(ws.title, table_prefix) or ws.title in {SUMMARY_SHEET_NAME, "summary"}]
    if len(delete_titles) >= len(wb.worksheets):
        wb.create_sheet("__一時退避__")
    for title in delete_titles:
        if title in wb.sheetnames:
            del wb[title]
    return min(table_index, len(wb.worksheets)), min(summary_index, len(wb.worksheets))


def save_workbook_atomic(wb, output_path: Path):
    output_path.parent.mkdir(parents=True, exist_ok=True)
    tmp = tempfile.NamedTemporaryFile(prefix=f".{output_path.stem}_", suffix=".tmp.xlsx", dir=output_path.parent, delete=False)
    tmp_name = Path(tmp.name)
    tmp.close()
    try:
        wb.save(tmp_name)
        os.replace(tmp_name, output_path)
    finally:
        try:
            tmp_name.unlink(missing_ok=True)
        except Exception:
            pass

def main():
    parser = argparse.ArgumentParser(description="現_打鍵前 / 現_打鍵後 CSV をテンプレート Excel 形式へ展開します。")
    parser.add_argument("--config", default=DEFAULT_CONFIG_FILE, help="設定ファイル JSON。既定: 設定ファイル.json")
    parser.add_argument("--template", default=None, help="添付テンプレート Excel")
    parser.add_argument("--template-sheet", default=None, help="テンプレート本体のシート名。既定: テーブル")
    parser.add_argument("--before", default=None, help="更新前CSVフォルダ")
    parser.add_argument("--after", default=None, help="更新後CSVフォルダ")
    parser.add_argument("--output", default=None, help="出力Excel")
    parser.add_argument("--sheet-name", default=None, help="出力シート名")
    parser.add_argument("--sql-name-mode", choices=["stem", "filename"], default=None, help="SQL の FROM に使う名前")
    parser.add_argument("--tables-per-sheet", type=int, default=None, help="1シートに出力するテーブル数。既定値は5。")
    args = parser.parse_args()
    logger, log_path = setup_logging(Path(__file__).stem)
    logger.info("処理開始")
    logger.info("ログファイル: %s", log_path)
    logger.info("設定ファイル: %s", args.config)

    settings = load_settings_file(Path(args.config) if args.config else None)
    common_cfg = settings.get("共通", {}) if isinstance(settings, dict) else {}
    old_cfg = settings.get("現CSV出力", {}) if isinstance(settings, dict) else {}

    # 新形式: パス + フォルダ名/ファイル名を分けず、1項目に完全パスを書く。
    # 旧形式も互換のため残す。新形式が設定されている場合は新形式を優先する。
    before_path_value = args.before or cfg_get(
        old_cfg, common_cfg,
        keys=["現_打鍵前CSVパス", "現_打鍵前CSVフォルダ", "現_打鍵前パス", "before_csv_path", "before_path"],
        default=None,
    )
    after_path_value = args.after or cfg_get(
        old_cfg, common_cfg,
        keys=["現_打鍵後CSVパス", "現_打鍵後CSVフォルダ", "現_打鍵後パス", "after_csv_path", "after_path"],
        default=None,
    )
    template_path_value = args.template or cfg_get(
        old_cfg, common_cfg,
        keys=["テンプレートファイルパス", "テンプレートExcelパス", "template_path"],
        default=None,
    )
    if template_path_value is None and not has_any_key(old_cfg, common_cfg, keys=["テンプレート入力パス", "テンプレートフォルダパス"]):
        template_path_value = cfg_get(old_cfg, keys=["テンプレートファイル", "テンプレート", "template"], default=None)

    output_path_value = args.output or cfg_get(
        old_cfg, common_cfg,
        keys=["出力ファイルパス", "出力Excelパス", "output_path"],
        default=None,
    )
    if output_path_value is None and not has_any_key(old_cfg, common_cfg, keys=["出力パス", "出力フォルダパス"]):
        output_path_value = cfg_get(old_cfg, keys=["出力ファイル", "output"], default=None)

    if before_path_value is not None:
        before_dir = resolve_path_list(before_path_value)
    else:
        before_base = cfg_get(old_cfg, common_cfg, keys=["現_打鍵前入力パス", "現_打鍵前CSV入力パス", "入力パス", "入力フォルダパス", "基準パス"], default=".")
        before_name = cfg_get(old_cfg, keys=["現_打鍵前フォルダ名", "before", "before_folder"], default="現_打鍵前")
        before_dir = resolve_folder_list(before_base, before_name)

    if after_path_value is not None:
        after_dir = resolve_path_list(after_path_value)
    else:
        after_base = cfg_get(old_cfg, common_cfg, keys=["現_打鍵後入力パス", "現_打鍵後CSV入力パス", "入力パス", "入力フォルダパス", "基準パス"], default=".")
        after_name = cfg_get(old_cfg, keys=["現_打鍵後フォルダ名", "after", "after_folder"], default="現_打鍵後")
        after_dir = resolve_folder_list(after_base, after_name)

    if template_path_value is not None:
        template_path = resolve_file_path(template_path_value, must_exist=True)
    else:
        template_base = cfg_get(old_cfg, common_cfg, keys=["テンプレート入力パス", "テンプレートパス", "テンプレートフォルダパス"], default=".")
        template_name = cfg_get(old_cfg, keys=["テンプレートファイル名", "テンプレート", "template"], default="テーブル.xlsx")
        template_path = join_if_relative(template_base, template_name)

    if output_path_value is not None:
        output_path = resolve_file_path(output_path_value, must_exist=any(ch in str(_first_ordered_value(output_path_value, "")) for ch in "*?["))
    else:
        output_base = cfg_get(old_cfg, common_cfg, keys=["出力パス", "出力フォルダパス"], default=".")
        output_name = cfg_get(old_cfg, keys=["出力ファイル名", "出力ファイル", "output"], default="比較結果.xlsx")
        output_path = resolve_wildcard_path(output_base, output_name, must_exist=any(ch in str(output_name) for ch in "*?["))

    template_sheet = args.template_sheet or cfg_get(old_cfg, keys=["テンプレートシート名", "template_sheet"], default="テーブル")
    sheet_name_arg = args.sheet_name or cfg_get(old_cfg, common_cfg, keys=["テーブルシート名", "sheet_name"], default="テーブル")
    sql_name_mode = args.sql_name_mode or cfg_get(old_cfg, keys=["SQL名モード", "sql_name_mode"], default="stem")
    tables_per_sheet = args.tables_per_sheet or int(cfg_get(old_cfg, common_cfg, keys=["1シートテーブル数", "tables_per_sheet"], default=5))
    if tables_per_sheet <= 0:
        raise ValueError("--tables-per-sheet は 1 以上を指定してください。")

    existing_action = normalize_existing_action(cfg_get(old_cfg, common_cfg, keys=["既存ファイル処理", "既存時動作", "existing_action"], default=ACTION_REPLACE_FILE))
    auto_width_enabled = bool_value(cfg_get(old_cfg, common_cfg, keys=["列幅自動調整", "自動列幅調整", "auto_fit_columns"], default=False))
    max_width = cfg_get(old_cfg, common_cfg, keys=["最大列幅", "max_column_width"], default=60)
    min_width = cfg_get(old_cfg, common_cfg, keys=["最小列幅", "min_column_width"], default=8)
    selector_enabled = bool_value(cfg_get(old_cfg, common_cfg, keys=["対象表選択有効", "対象表選択を使用", "enable_table_selector"], default=False))
    selector_sheet = cfg_get(old_cfg, common_cfg, keys=["対象表選択シート名", "選択シート名", "table_selector_sheet"], default=None)
    selector_mark_col = cfg_get(old_cfg, common_cfg, keys=["対象表選択判定列", "選択判定列", "table_selector_mark_column"], default=None)
    selector_table_col = cfg_get(old_cfg, common_cfg, keys=["対象表選択表名列", "選択表名列", "table_selector_table_column"], default=None)
    selector_legacy_condition_col = cfg_get(old_cfg, common_cfg, keys=["対象表選択条件列", "選択条件列", "table_selector_condition_column"], default=None)
    selector_before_condition_col = cfg_get(old_cfg, common_cfg, keys=["対象表選択打鍵前条件列", "打鍵前条件列", "table_selector_before_condition_column"], default=selector_legacy_condition_col)
    selector_after_condition_col = cfg_get(old_cfg, common_cfg, keys=["対象表選択打鍵後条件列", "打鍵後条件列", "table_selector_after_condition_column"], default=selector_legacy_condition_col)
    selector_before_current_condition_col = cfg_get(old_cfg, common_cfg, keys=["対象表選択打鍵前現条件列", "打鍵前現条件列", "table_selector_before_current_condition_column"], default=selector_before_condition_col)
    selector_before_new_condition_col = cfg_get(old_cfg, common_cfg, keys=["対象表選択打鍵前新条件列", "打鍵前新条件列", "table_selector_before_new_condition_column"], default=selector_before_condition_col)
    selector_after_current_condition_col = cfg_get(old_cfg, common_cfg, keys=["対象表選択打鍵後現条件列", "打鍵後現条件列", "table_selector_after_current_condition_column"], default=selector_after_condition_col)
    selector_after_new_condition_col = cfg_get(old_cfg, common_cfg, keys=["対象表選択打鍵後新条件列", "打鍵後新条件列", "table_selector_after_new_condition_column"], default=selector_after_condition_col)
    selector_mark_value = cfg_get(old_cfg, common_cfg, keys=["対象表選択判定値", "選択判定値", "table_selector_mark_value"], default="〇")

    if not template_path.exists():
        raise FileNotFoundError(f"テンプレート Excel が存在しません: {template_path}")
    if output_path.exists() and existing_action == ACTION_ERROR:
        raise FileExistsError(f"出力ファイルが既に存在します: {output_path}")

    logger.info("現_打鍵前CSVパス: %s", [str(p) for p in before_dir])
    logger.info("現_打鍵後CSVパス: %s", [str(p) for p in after_dir])
    logger.info("テンプレートファイル: %s", template_path)
    logger.info("出力ファイル: %s", output_path)
    logger.info("既存ファイル処理: %s", existing_action)
    logger.info("対象表選択有効: %s", selector_enabled)

    selected_table_names = None
    selected_table_conditions = {}
    if selector_enabled:
        selected_table_names = read_selected_table_names_from_workbook(
            output_path,
            selector_sheet,
            selector_mark_col,
            selector_table_col,
            selector_mark_value,
        )
        selected_table_conditions = read_selected_table_conditions_from_workbook(
            output_path,
            selector_sheet,
            selector_mark_col,
            selector_table_col,
            selector_before_current_condition_col,
            None,
            selector_after_current_condition_col,
            None,
            selector_mark_value,
        )
        logger.info(
            "対象表選択シート: %s / 判定列: %s / 表名列: %s / 打鍵前現条件列: %s / 打鍵前新条件列: %s / 打鍵後現条件列: %s / 打鍵後新条件列: %s / 判定値: %s",
            selector_sheet,
            selector_mark_col,
            selector_table_col,
            selector_before_current_condition_col,
            selector_before_new_condition_col,
            selector_after_current_condition_col,
            selector_after_new_condition_col,
            selector_mark_value,
        )
        non_empty_current_condition_count = sum(
            1
            for conditions in selected_table_conditions.values()
            for key, value in conditions.items()
            if key in {"before_current", "after_current"} and str(value or "").strip()
        )
        logger.info("対象表選択件数: %s / 現抽出条件対象表数: %s / 非空現抽出条件数: %s", 0 if selected_table_names is None else len(selected_table_names), len(selected_table_conditions), non_empty_current_condition_count)

    template_wb = load_workbook(template_path)
    compare_exclusions = read_compare_exclusions(template_wb)
    if template_sheet in template_wb.sheetnames:
        tws = template_wb[template_sheet]
    else:
        template_candidates = [ws for ws in template_wb.worksheets if ws.title != COMPARE_EXCLUDE_SHEET_NAME]
        tws = template_candidates[0] if template_candidates else template_wb.active
    layout = template_layout(tws)

    records = build_table_records(csv_files(before_dir), csv_files(after_dir), sql_name_mode)
    data_table_aliases = table_aliases_with_csv_data_from_records(records)
    selector_note_by_alias = selector_notes_by_table_alias(records)
    logger.info("CSV突合テーブル数: %s", len(records))
    logger.info("CSVデータありテーブル数(alias含む): %s", len(data_table_aliases))
    if selected_table_names is not None:
        records = [rec for rec in records if record_matches_selected_tables(rec, selected_table_names)]
        apply_selected_table_conditions(records, selected_table_conditions)
        logger.info("対象表選択後テーブル数: %s", len(records))
    if not records:
        raise RuntimeError("CSVファイルが見つかりません。")
    output_records = assign_output_positions(records, sheet_name_arg, tables_per_sheet)

    replace_sheets_mode = output_path.exists() and existing_action == ACTION_REPLACE_SHEETS
    preserved_image_sheet_parts = {}
    preserved_image_sheet_source = None
    preserved_image_sheet_indexes = {}
    if replace_sheets_mode:
        try:
            existing_sheet_names = set(sheet_xml_paths(output_path).keys())
        except Exception:
            existing_sheet_names = set()
        existing_sheet_indexes = sheet_index_map(output_path)
        image_sheet_names = sheet_names_with_image_parts(output_path)
        preserved_image_sheet_names = {
            name for name in existing_sheet_names
            if not is_table_output_sheet(name, sheet_name_arg) and name not in {SUMMARY_SHEET_NAME, "summary"}
            and name in image_sheet_names
        }
        preserved_image_sheet_indexes = {
            name: existing_sheet_indexes[name]
            for name in preserved_image_sheet_names
            if name in existing_sheet_indexes
        }
        preserved_image_sheet_parts = collect_sheet_parts_to_preserve(output_path, preserved_image_sheet_names)
        if preserved_image_sheet_parts:
            tmp_preserve = tempfile.NamedTemporaryFile(prefix=f".{output_path.stem}_image_sheets_", suffix=".xlsx", dir=output_path.parent, delete=False)
            preserved_image_sheet_source = Path(tmp_preserve.name)
            tmp_preserve.close()
            shutil.copy2(output_path, preserved_image_sheet_source)
        wb = load_workbook(output_path)
        table_insert_index, summary_insert_index = remove_existing_output_sheets(wb, sheet_name_arg)
        # テーブル検証シートがテーブルシートより前に戻される場合、
        # 後でテーブル検証を挿入するとテーブル位置が右へずれるため、先に1つ左へ補正する。
        adjusted_table_insert_index = max(0, table_insert_index - 1) if summary_insert_index <= table_insert_index else table_insert_index
    else:
        wb = template_wb
        table_insert_index = None
        adjusted_table_insert_index = None
        summary_insert_index = 0

    original_sheet_title = tws.title
    table_sheet_targets = {}
    table_sheets_to_rename = []

    sheet_names = sorted(set(r.output_sheet for r in output_records), key=lambda s: (len(s), s))
    for sheet_name in sheet_names:
        tmp_name = unique_sheet_name(wb, f"__{sheet_name}__")
        create_index = None if adjusted_table_insert_index is None else adjusted_table_insert_index + len(table_sheets_to_rename)
        ws = wb.create_sheet(tmp_name, create_index)
        copy_sheet_page_settings(tws, ws)
        sheet_records = [r for r in output_records if r.output_sheet == sheet_name]
        row = 1
        max_col_used = TEMPLATE_COLS
        header_rows = set()
        cf_targets = []
        for i, rec in enumerate(sheet_records):
            row, block_last_col = write_block(tws, ws, row, rec, layout, header_rows, cf_targets, compare_exclusions)
            max_col_used = max(max_col_used, block_last_col)
            if i != len(sheet_records) - 1:
                row += 1
        last_row = max(1, row - 1)
        normalize_fonts_and_row_heights(ws, last_row, max_col_used, header_rows)
        table_sheets_to_rename.append((ws, sheet_name))
        table_sheet_targets[sheet_name] = cf_targets

    write_summary_sheet(wb, records, summary_insert_index)

    if not replace_sheets_mode:
        remove_sheet_titles = {original_sheet_title, COMPARE_EXCLUDE_SHEET_NAME}
        for ws in list(wb.worksheets):
            if ws.title in remove_sheet_titles:
                wb.remove(ws)
    else:
        if "__一時退避__" in wb.sheetnames and len(wb.worksheets) > 1:
            del wb["__一時退避__"]

    for ws, sheet_name in table_sheets_to_rename:
        ws.title = sheet_name

    if selector_enabled:
        mismatch_count = mark_selector_csv_mismatches(
            wb,
            selector_sheet,
            selector_mark_col,
            selector_table_col,
            selector_mark_value,
            data_table_aliases,
            selected_table_names,
            selector_note_by_alias,
        )
        logger.info("対象表選択CSV不一致件数: %s", mismatch_count)

    auto_fit_column_widths(
        wb,
        enabled=auto_width_enabled,
        max_width=max_width,
        min_width=min_width,
        sheet_names=[SUMMARY_SHEET_NAME] + [s for s in table_sheet_targets.keys()],
    )

    try:
        wb.calculation.fullCalcOnLoad = True
        wb.calculation.forceFullCalc = True
    except Exception:
        pass

    save_workbook_atomic(wb, output_path)
    inject_x14_conditional_formatting(output_path, template_path, table_sheet_targets, layout, tws.title)
    try:
        if preserved_image_sheet_source and preserved_image_sheet_parts:
            reinject_preserved_sheet_parts(output_path, preserved_image_sheet_source, preserved_image_sheet_parts)
            restore_sheet_order_positions(output_path, preserved_image_sheet_indexes)
    finally:
        if preserved_image_sheet_source:
            try:
                preserved_image_sheet_source.unlink(missing_ok=True)
            except Exception:
                pass
    print_and_log(f"出力しました: {output_path} / summary={SUMMARY_SHEET_NAME} / mode={existing_action}", logger)
    print_and_log(f"ログファイル: {log_path}", logger)
    if selected_table_names is not None:
        print_and_log(f"対象表選択: {len(selected_table_names)}件指定 / {len(records)}件出力", logger)
    logger.info("処理終了")


if __name__ == "__main__":
    try:
        main()
    except Exception:
        logging.exception("処理中にエラーが発生しました。")
        raise
